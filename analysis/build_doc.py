import re
import sys
from docx import Document
from docx.enum.text import WD_COLOR_INDEX, WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Pt, Cm
from docx.oxml.ns import qn
from stats_core import *

OUT = sys.argv[1]

# ---------------------------------------------------------------- references
# verified=True  -> citation copied from the reference list of Lee et al. (AJSM 2022) PDF provided by the author
# verified=False -> added from general knowledge; original must be checked by the author (highlighted in the list)
REFS = {
    'lee2022': (True, 'Lee SH, Yang JH, Kim I. Nonanatomic all-inside arthroscopic anterior talofibular ligament repair with a high-position anchor versus anatomic repair: an analysis based on 3D CT. Am J Sports Med. 2022;50(8):2134-2144. doi:10.1177/03635465221097119'),
    'brostrom': (True, 'Broström L. Sprained ankles: V. Treatment and prognosis in recent ligament ruptures. Acta Chir Scand. 1966;132:537-550.'),
    'gould': (False, 'Gould N, Seligson D, Gassman J. Early and late repair of lateral ligament of the ankle. Foot Ankle. 1980;1(2):84-89.'),
    'vega2013': (True, 'Vega J, Golanó P, Pellegrino A, Rabat E, Peña F. All-inside arthroscopic lateral collateral ligament repair for ankle instability with a knotless suture anchor technique. Foot Ankle Int. 2013;34:1701-1709.'),
    'takao2016': (True, 'Takao M, Matsui K, Stone JW, et al. Arthroscopic anterior talofibular ligament repair for lateral instability of the ankle. Knee Surg Sports Traumatol Arthrosc. 2016;24:1003-1006.'),
    'matsui2016': (True, 'Matsui K, Takao M, Miyamoto W, Matsushita T. Early recovery after arthroscopic repair compared to open repair of the anterior talofibular ligament for lateral instability of the ankle. Arch Orthop Trauma Surg. 2016;136:93-100.'),
    'matsui2014': (True, 'Matsui K, Takao M, Miyamoto W, Innami K, Matsushita T. Arthroscopic Broström repair with Gould augmentation via an accessory anterolateral port for lateral instability of the ankle. Arch Orthop Trauma Surg. 2014;134:1461-1467.'),
    'vega2020': (True, 'Vega J, Malagelada F, Dalmau-Pastor M. Arthroscopic all-inside ATFL and CFL repair is feasible and provides excellent results in patients with chronic ankle instability. Knee Surg Sports Traumatol Arthrosc. 2020;28:116-123.'),
    'drakos2014': (True, 'Drakos MC, Behrens SB, Paller D, Murphy C, DiGiovanni CW. Biomechanical comparison of an open vs arthroscopic approach for lateral ankle instability. Foot Ankle Int. 2014;35:809-815.'),
    'krips2000': (True, 'Krips R, van Dijk CN, Halasi T, et al. Anatomical reconstruction versus tenodesis for the treatment of chronic anterolateral instability of the ankle joint: a 2- to 10-year follow-up, multicenter study. Knee Surg Sports Traumatol Arthrosc. 2000;8:173-179.'),
    'shoji2019': (True, 'Shoji H, Teramoto A, Sakakibara Y, et al. Kinematics and laxity of the ankle joint in anatomic and nonanatomic anterior talofibular ligament repair: a biomechanical cadaveric study. Am J Sports Med. 2019;47:667-673.'),
    'caputo2009': (True, 'Caputo AM, Lee JY, Spritzer CE, et al. In vivo kinematics of the tibiotalar joint after lateral ankle instability. Am J Sports Med. 2009;37:2241-2248.'),
    'burks1994': (True, 'Burks RT, Morgan J. Anatomy of the lateral ankle ligaments. Am J Sports Med. 1994;22:72-77.'),
    'kakegawa2019': (True, 'Kakegawa A, Mori Y, Tsuchiya A, Sumitomo N, Fukushima N, Moriizumi T. Independent attachment of lateral ankle ligaments: anterior talofibular and calcaneofibular ligaments—a cadaveric study. J Foot Ankle Surg. 2019;58:717-722.'),
    'matsui2017': (True, 'Matsui K, Oliva XM, Takao M, et al. Bony landmarks available for minimally invasive lateral ankle stabilization surgery: a cadaveric anatomical study. Knee Surg Sports Traumatol Arthrosc. 2017;25:1916-1924.'),
    'nakasa2021': (True, 'Nakasa T, Ikuta Y, Ota Y, et al. Safe angles of ATFL and CFL anchor insertion into anatomical attachment of fibula in a lateral ankle ligament repair. J Orthop Sci. 2021;26:156-161. doi:10.1016/j.jos.2020.02.011'),
    'teramoto2018': (True, 'Teramoto A, Shoji H, Sakakibara Y, Suzuki T, Watanabe K, Yamashita T. The distal margin of the lateral malleolus visible under ankle arthroscopy (articular tip) from the anteromedial portal, is separate from the ATFL attachment site of the fibula: a cadaver study. J Orthop Sci. 2018;23:565-569.'),
    'thes2016': (True, 'Thès A, Klouche S, Ferrand M, Hardy P, Bauer T. Assessment of the feasibility of arthroscopic visualization of the lateral ligament of the ankle: a cadaveric study. Knee Surg Sports Traumatol Arthrosc. 2016;24:985-990.'),
    'lee2021': (True, 'Lee SH, Cho HG, Yang JH. Additional inferior extensor retinaculum augmentation after all-inside arthroscopic anterior talofibular ligament repair for chronic ankle instability is not necessary. Am J Sports Med. 2021;49:1721-1731.'),
    'leeyang2021': (True, 'Lee SH, Yang JH. All-inside arthroscopic anatomic anterior talofibular ligament repair for anterolateral ankle instability using a knotless suture anchor, allowing for tension adjustment. Arthrosc Tech. 2021;10:e925-e929. doi:10.1016/j.eats.2020.11.013'),
    'beighton': (True, 'Beighton P, Solomon L, Soskolne CL. Articular mobility in an African population. Ann Rheum Dis. 1973;32:413-418.'),
    'karlsson1991': (False, 'Karlsson J, Peterson L. Evaluation of ankle joint function: the use of a scoring scale. Foot. 1991;1(1):15-19.'),
    'kitaoka1994': (False, 'Kitaoka HB, Alexander IJ, Adelaar RS, Nunley JA, Myerson MS, Sanders M. Clinical rating systems for the ankle-hindfoot, midfoot, hallux, and lesser toes. Foot Ankle Int. 1994;15(7):349-353.'),
    'tegner1985': (False, 'Tegner Y, Lysholm J. Rating systems in the evaluation of knee ligament injuries. Clin Orthop Relat Res. 1985;(198):43-49.'),
    'budiman1991': (False, 'Budiman-Mak E, Conrad KJ, Roach KE. The Foot Function Index: a measure of foot pain and disability. J Clin Epidemiol. 1991;44(6):561-570.'),
    'roos2001': (True, 'Roos EM, Brandsson S, Karlsson J. Validation of the Foot and Ankle Outcome Score for ankle ligament reconstruction. Foot Ankle Int. 2001;22:788-794.'),
    'martin2005': (False, 'Martin RL, Irrgang JJ, Burdett RG, Conti SF, Van Swearingen JM. Evidence of validity for the Foot and Ankle Ability Measure (FAAM). Foot Ankle Int. 2005;26(11):968-983.'),
    'herdman2011': (False, 'Herdman M, Gudex C, Lloyd A, et al. Development and preliminary testing of the new five-level version of EQ-5D (EQ-5D-5L). Qual Life Res. 2011;20(10):1727-1736.'),
}
ORDER = []


def cite(keys):
    nums = []
    for k in keys.split(','):
        k = k.strip()
        if k not in REFS:
            raise KeyError(k)
        if k not in ORDER:
            ORDER.append(k)
        nums.append(ORDER.index(k) + 1)
    nums = sorted(set(nums))
    # compress ranges
    parts, i = [], 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        parts.append(str(nums[i]) if j == i else (f'{nums[i]},{nums[j]}' if j == i + 1 else f'{nums[i]}-{nums[j]}'))
        i = j + 1
    return '[' + ','.join(parts) + ']'


# ---------------------------------------------------------------- docx helpers
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
for side in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
    setattr(sec, side, Cm(2.5))
st = doc.styles['Normal']
st.font.name = 'Times New Roman'
st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.5
for lvl in (1, 2, 3):
    hs = doc.styles[f'Heading {lvl}']
    hs.font.name = 'Times New Roman'
    hs.element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
    hs.font.color.rgb = None
    hs.font.size = Pt({1: 14, 2: 12, 3: 11}[lvl])

TOKEN = re.compile(r'(\[\[.*?\]\]|\{\{.*?\}\}|<<[^>]+>>|\*\*.*?\*\*)')


def add_runs(p, text, size=None, bold=False, italic=False):
    """[[x]] = yellow (author input needed); {{x}} = turquoise (changed vs. abstract / check);
    <<key,key>> = citation; **x** = bold."""
    for part in TOKEN.split(text):
        if not part:
            continue
        hl = None
        b = bold
        if part.startswith('[[') and part.endswith(']]'):
            part, hl = '[' + part[2:-2] + ']', WD_COLOR_INDEX.YELLOW
        elif part.startswith('{{') and part.endswith('}}'):
            part, hl = part[2:-2].replace('[[', '[').replace(']]', ']'), WD_COLOR_INDEX.TURQUOISE
            part = re.sub(r'<<([^>]+)>>', lambda m: cite(m.group(1)), part)
        elif part.startswith('<<') and part.endswith('>>'):
            part = cite(part[2:-2])
        elif part.startswith('**') and part.endswith('**'):
            part, b = part[2:-2], True
        r = p.add_run(part)
        r.bold = b
        r.italic = italic
        if size:
            r.font.size = Pt(size)
        if hl:
            r.font.highlight_color = hl
    return p


def para(text, style=None, align=None, size=None, bold=False, italic=False, space_after=None, line=None):
    p = doc.add_paragraph(style=style)
    add_runs(p, text, size=size, bold=bold, italic=italic)
    if align:
        p.alignment = align
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    if line is not None:
        p.paragraph_format.line_spacing = line
    return p


def h(text, lvl=1):
    return doc.add_heading(text, level=lvl)


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def table(header, rows, widths, note=None, title=None):
    if title:
        para(title, bold=True, size=10, space_after=3, line=1.0)
    t = doc.add_table(rows=1, cols=len(header))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, txt in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ''
        add_runs(c.paragraphs[0], txt, size=8.5, bold=True)
    for row in rows:
        cells = t.add_row().cells
        for i, txt in enumerate(row):
            cells[i].text = ''
            add_runs(cells[i].paragraphs[0], str(txt), size=8.5, bold=(i == 0 and str(txt).isupper()))
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
            for p in row.cells[i].paragraphs:
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_after = Pt(0)
                if i > 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if note:
        para(note, size=8.5, line=1.0, space_after=12)
    return t


# ---------------------------------------------------------------- numbers
N = len(S)
NPAT = S.id.nunique()
ng = {g: int((S.G == g).sum()) for g in GROUPS}
nx = {g: int((XR.G == g).sum()) for g in GROUPS}
teg_up = int((S.d_teg > 0).sum())
teg_same = int((S.d_teg == 0).sum())
teg_down = int((S.d_teg < 0).sum())
teg_keep = teg_up + teg_same
pct = lambda a, b: f'{a / b * 100:.1f}'
ffi_kw = kw(S, 'ffi')
ffi_ph = mwu_pairs(S, 'ffi')
ffi_rho = rho(S, 'ffi')
fem_pat = int((S.drop_duplicates('id').sex == 'F').sum())
male_pat = NPAT - fem_pat
right = int((S.side == 1).sum())
left = int((S.side == 2).sum())
XRc = XR[XR.c_tt.notna()]
opmin, opmax = S.opdate.min(), S.opdate.max()

# ================================================================= AUTHOR QUERIES
para('ArthroBrostrom 3D-CT anchor position — manuscript draft v0.1', bold=True, size=14, line=1.0)
para('Draft generated 2026-09-26 · for discussion between co-authors · not for submission', italic=True, size=9, line=1.0)
para('How to read the highlights', bold=True, line=1.0)
para('[[Yellow = information only the authors have (not written by the draft author; please fill in)]]', size=10, line=1.0)
para('{{Turquoise = a value or statement that differs from the congress abstract, or an interpretation that needs your confirmation}}', size=10, line=1.0)
para('In the reference list, highlighted entries were added from general knowledge and have NOT been checked against an original; please send the PDF or confirm. Unhighlighted entries were taken verbatim from the reference list of Lee et al. (AJSM 2022). Every number in the tables is computed directly from “update ver2.5.xlsx” by script; no value was typed by hand.', size=10, line=1.0)

para('저자 확인 요청 사항 (Author queries)', bold=True, size=12, line=1.0)
queries = [
    ('Cohort / 81 ankles', f'81례 = Sheet1에서 CT anchor 비율이 있는 증례 중 2025년 이전 수술례(No.1–94 중 CT/설문 결측 13례 제외). 동일 등록번호 환자(No.72, No.80)가 양측 수술 → 80명 81례로 확인됨. 이 정의가 맞는지 확인 부탁드립니다. 또한 전체 수술 환자 수와 제외 사유별 인원(flow chart용)이 필요합니다.'),
    ('2-year follow-up', f'No.94 (수술일 2024-12-31)는 오늘(2026-09-26) 기준 약 21개월로, “최소 2년 추시” 조건을 만족하지 않습니다. 최종 설문을 언제 받았는지, 포함 유지 여부를 알려주세요. 또한 No.93(2024-09-24)은 정확히 2년입니다.'),
    ('Follow-up period', '엑셀의 “Last F/U” 날짜는 외래 방문일로 보이며 수술 후 수개월인 경우가 많아, 최종 설문(PROM) 시행일로 보기 어렵습니다. 각 증례의 최종 설문 시행일(또는 평균/범위 추시기간)과 설문 방법(외래/전화/우편)을 알려주세요.'),
    ('Tegner (변경됨)', f'현재 데이터: Tegner {ms(S.teg_pre)} → {ms(S.teg_post)}, Wilcoxon P = {fp(wil(S, "teg_pre", "teg_post"))}. 유지/향상 {teg_keep}례 ({pct(teg_keep, N)}%): 향상 {teg_up}, 유지 {teg_same}, 감소 {teg_down}. 초록의 “68예 (87.7%)”는 68/81=84.0%로 계산이 맞지 않아, 71/81=87.7%가 맞는 것으로 보입니다.'),
    ('FFI 사후분석 (중요)', f'초록에는 “다중비교 보정 후 유의하지 않음”으로 되어 있으나, 현재 데이터로 Kruskal–Wallis P = {fp(ffi_kw)}, 사후 Mann–Whitney 해부학적 vs 준해부학적 P = {fp(ffi_ph["AS"])} (Bonferroni 기준 P < .017에서도 유의), Dunn–Bonferroni 보정 P = .035로 여전히 유의합니다. 초록에서 사용한 보정 방법(예: 전체 결과변수 수로 보정)을 알려주시면 그에 맞춰 수정하겠습니다. 본문은 현재 계산값 그대로 기술했습니다.'),
    ('“3차원 CT 분석”의 의미', f'초록의 “3차원 CT 분석에서는 유의하지 않았다”를 anchor 비율(연속변수)과 FFI의 Spearman 상관(ρ = {ffi_rho[0]:.2f}, P = {fp(ffi_rho[1])})으로 해석했습니다. 다른 분석이었다면 알려주세요.'),
    ('통계 방법/프로그램', '본 초안은 Lee et al.과 동일한 비모수 방법(Kruskal–Wallis, Wilcoxon signed-rank, Mann–Whitney U post hoc with P < .017, χ² test)과 Spearman 상관을 Python (SciPy 1.17)으로 계산했습니다. 초록 수치(VAS/KP/AOFAS)는 정확히 재현되었습니다. 실제 사용한 프로그램(SPSS 등)과 검정법을 알려주세요. 사전 검정력 분석(power analysis) 여부도 필요합니다.'),
    ('수술 방법', '수술 기법 전체가 필요합니다: 마취/체위, portal, arthroscope, anchor 종류·크기·개수(엑셀 메모에 “Arthrex anchor”, “IER reinforce” 언급), ATFL 봉합 방법, IER(Gould) 보강 여부(전례 시행?), CFL 처리, 동반 술식(OLT, impingement 등), 집도의 수(단일 술자?).'),
    ('재활', '술후 고정 기간, 체중부하 시작 시점, 운동 복귀 시점 등 재활 프로토콜.'),
    ('CT 측정', 'CT 촬영 시점(술후 며칠), 장비, 3D 재구성 소프트웨어, 측정자(몇 명, 맹검 여부), 측정 신뢰도(ICC/kappa).'),
    ('스트레스 방사선', '장비(Telos 등)·하중(N)·측정 방법, 술전 촬영 시점, 최종 촬영 시점. 참고: 술전 전방전위가 건측과 거의 차이가 없습니다(환측−건측 평균 0.52 mm). 측정 방법 확인이 필요합니다. 최종 스트레스 사진이 없는 9례(각 군 2/2/5)가 있습니다.'),
    ('점수 정의', 'FAOS 총점 = 5개 하위척도의 평균으로 계산되어 있습니다(원래 FAOS는 총점을 정의하지 않음). No.51은 FAOS 총점만 있고 하위척도가 없습니다. EQ-5D-5L은 5–11의 값으로, 효용지수가 아닌 5개 차원 수준의 합(5 = 최상)으로 보입니다 → 한국형 가치세트로 index 변환할지 결정 필요. Karlsson-Peterson/AOFAS/FFI/FAAM은 한국어 번역판 사용 여부와 해당 인용문헌.'),
    ('포함/제외 기준, IRB', '포함/제외 기준, IRB 승인번호, 동의서 면제 여부, 연구기관명(경희대학교병원/강동경희대병원 등), 저자 영문 이름.'),
    ('대상 학술지', '투고 예정 학술지(영문/국문, 분량·참고문헌 양식)를 알려주시면 형식을 맞추겠습니다. 현재는 영문 AJSM/FAI 스타일(Vancouver 번호식)로 작성했습니다.'),
    ('필요 문헌 원문', '아래 참고문헌 중 하이라이트된 항목의 원문 PDF, 그리고 가능하면: (1) Lee et al. 2022 이후 anchor 위치/비해부학적 봉합 관련 임상 연구, (2) 관절경적 변형 브로스트롬(IER 보강 포함) 장기 추시 연구, (3) 한국어판 FAOS/FAAM/FFI/EQ-5D 타당도 논문, (4) Karlsson score 원저. 이 원문을 받으면 Discussion을 보강하겠습니다.'),
]
table(['#', 'Item', '내용'], [(str(i + 1), a, b) for i, (a, b) in enumerate(queries)], [0.8, 3.2, 12])
page_break()

# ================================================================= TITLE PAGE
para('Original Article', italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)
para('Clinical and Radiologic Outcomes of Arthroscopic Modified Broström Procedure According to Fibular Anchor Position: A 3-Dimensional Computed Tomography–Based Analysis With a Minimum 2-Year Follow-up',
     bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER)
para('Running title: Anchor position in arthroscopic modified Broström procedure', align=WD_ALIGN_PARAGRAPH.CENTER, size=10)
para('[[English names]]: 이정현, 서동욱, 경민규, 정비오*', align=WD_ALIGN_PARAGRAPH.CENTER)
para('Department of Orthopaedic Surgery, Kyung Hee University School of Medicine, Seoul, Korea [[hospital name / address to confirm]]', align=WD_ALIGN_PARAGRAPH.CENTER, size=10)
para('*Corresponding author: 정비오 [[English name, address, phone]]; e-mail: biojeong@gmail.com', align=WD_ALIGN_PARAGRAPH.CENTER, size=10)
para('Conflict of interest: [[ ]]  ·  Funding: [[ ]]  ·  IRB approval: [[institution, number]]', align=WD_ALIGN_PARAGRAPH.CENTER, size=10)
page_break()

# ================================================================= ABSTRACT
h('Abstract')
para('**Background:** Anatomic placement of the fibular anchor is considered important in arthroscopic lateral ankle ligament repair. A previous 3-dimensional computed tomography (3D CT) study reported inferior outcomes when the anchor was placed in a nonanatomic (high) position after all-inside arthroscopic anterior talofibular ligament (ATFL) repair. This study evaluated whether clinical and radiologic outcomes of the arthroscopic modified Broström procedure differ according to fibular anchor position.')
para(f'**Methods:** {N} ankles in {NPAT} patients who underwent the arthroscopic modified Broström procedure with a minimum follow-up of 2 years{{{{ [[confirm; see query 2]]}}}} were retrospectively reviewed. On postoperative 3D CT, the distance from the anchor center to the fibular obscure tubercle (FOT) was divided by the distance between the fibular anterior tubercle (FAT) and the FOT. Ankles were classified as anatomic (<25%, n = {ng["A"]}), subanatomic (25% to <50%, n = {ng["S"]}), or nonanatomic (≥50%, n = {ng["N"]}). The visual analog scale (VAS) for pain, Karlsson-Peterson (KP) score, American Orthopaedic Foot and Ankle Society (AOFAS) score, and Tegner activity score were compared between the preoperative and final follow-up evaluations; the Foot Function Index (FFI), Foot and Ankle Outcome Score (FAOS), EQ-5D-5L, and Foot and Ankle Ability Measure (FAAM) were assessed at the final follow-up. Varus talar tilt, anterior talar translation, and side-to-side differences were measured on stress radiographs.')
para(f'**Results:** The VAS decreased from {ms(S.pre_vas)} to {ms(S.vas)} (P {fp(wil(S,"pre_vas","vas"))}), the KP score improved from {ms(S.pre_kp)} to {ms(S.kp)}, and the AOFAS score improved from {ms(S.pre_aofas)} to {ms(S.aofas)} (both P {fp(wil(S,"pre_aofas","aofas"))}). The Tegner activity score increased from {{{{{ms(S.teg_pre)} to {ms(S.teg_post)} (P = {fp(wil(S,"teg_pre","teg_post"))})}}}}, and the activity level was maintained or improved in {{{{{teg_keep} ankles ({pct(teg_keep,N)}%)}}}}. Final scores and pre- to postoperative changes in the VAS, KP, AOFAS, and Tegner scores did not differ among the 3 groups (all P > .05). The FFI differed among the groups (P = {fp(ffi_kw)}), being lower in the subanatomic group than in the anatomic group{{{{ (post hoc P = {fp(ffi_ph["AS"])}) [[see query 5]]}}}}; however, the FFI was not correlated with the anchor position ratio (ρ = {ffi_rho[0]:.2f}, P = {fp(ffi_rho[1])}). The FAOS, FAAM, and EQ-5D-5L did not differ among the groups. Final talar tilt, anterior translation, side-to-side differences, and radiologic changes did not differ among the groups (all P > .05).')
para('**Conclusion:** No significant differences in clinical or radiologic outcomes were found according to fibular anchor position after the arthroscopic modified Broström procedure. Pain and functional scores improved significantly in the overall cohort, and the activity level was maintained or improved in most patients.')
para('**Level of Evidence:** Level III, retrospective cohort study [[confirm]]')
para('**Keywords:** anterior talofibular ligament; chronic ankle instability; arthroscopic modified Broström procedure; suture anchor position; 3-dimensional computed tomography')
page_break()

# ================================================================= INTRODUCTION
h('Introduction')
para('Direct anatomic repair of the lateral ankle ligaments, originally described by Broström, remains the favored surgical procedure for chronic ankle instability (CAI) <<brostrom,lee2022>>, and augmentation with the inferior extensor retinaculum (IER) has been added as a modification of this procedure <<gould,matsui2014>>. With advances in arthroscopic techniques, all-inside arthroscopic ATFL repair and the arthroscopic Broström procedure have emerged as effective alternatives to the open procedure <<vega2013,takao2016,matsui2016,vega2020>>. A biomechanical study showed comparable strength between open and arthroscopic Broström repairs <<drakos2014>>, and clinical studies have reported favorable outcomes with early recovery after arthroscopic repair <<matsui2016,vega2020>>.')
para('Restoration of normal anatomy is regarded as an important factor for a good prognosis after lateral ankle ligament surgery <<krips2000,lee2022>>. In a cadaveric study, Shoji et al <<shoji2019>> found that ATFL repair through a fibular bone tunnel located proximal to the anatomic footprint (nonanatomic repair) resulted in increased inversion and internal rotation kinematics and internal rotation laxity compared with the intact state, resembling the kinematics of the ATFL-deficient ankle <<caputo2009,shoji2019>>. However, arthroscopic identification of the fibular ATFL footprint is not always straightforward. One cadaveric study showed that only the portion 7 to 10 mm proximal to the tip of the lateral malleolus is visible through standard arthroscopic portals, raising the possibility of unintended proximal anchor placement <<teramoto2018>>, whereas another study reported that arthroscopic identification of the lateral ligaments and their footprints is feasible <<thes2016>>. Bony landmarks such as the FOT and FAT have been proposed to help identify the ligament footprints during minimally invasive surgery <<matsui2017,nakasa2021>>.')
para('Using postoperative 3D CT, Lee et al <<lee2022>> classified the fibular anchor position after all-inside arthroscopic ATFL repair according to its relative location between the FAT and FOT and reported that nonanatomic repair with a high-position anchor resulted in inferior FAOS, Karlsson scores, and posturographic fall risk compared with anatomic repair, although stress radiographic outcomes did not differ among groups. Their procedure consisted of ATFL repair with a single knotless suture anchor <<lee2022,leeyang2021>>. Whether anchor position has a similar influence on outcomes after the arthroscopic modified Broström procedure [[with IER augmentation — confirm technique]], and when a broader set of patient-reported outcome measures is used, remains unclear.')
para('Therefore, the purpose of this study was to compare the clinical and radiologic outcomes of the arthroscopic modified Broström procedure according to the fibular anchor position measured on 3D CT, using the classification described by Lee et al <<lee2022>>. We hypothesized that outcomes would differ according to anchor position, as previously reported.')

# ================================================================= METHODS
h('Materials and Methods')
h('Study Design and Patients', 2)
para(f'This retrospective study was approved by the institutional review board of [[institution]] ([[IRB number]]), and the requirement for informed consent was [[waived / obtained]]. We reviewed patients who underwent the arthroscopic modified Broström procedure for CAI between {opmin:%B %Y} and {opmax:%B %Y} [[performed by a single surgeon? — confirm]]. The inclusion criteria were as follows: (1) CAI with failure of nonoperative treatment for at least [[ ]] months, (2) the arthroscopic modified Broström procedure, (3) available postoperative 3D CT, and (4) a minimum of 2 years of clinical follow-up. The exclusion criteria were [[e.g., previous surgery on the affected ankle, arthritis, concomitant procedures such as ligament reconstruction or osteotomy, neuromuscular disorder — please provide]]. Of [[total number]] ankles, [[ ]] were excluded ([[reasons]]), and {N} ankles in {NPAT} patients (one patient underwent bilateral surgery) were included in the final analysis (Figure 1).')

h('Surgical Technique', 2)
para('[[To be provided by the authors: anesthesia and positioning; portals and arthroscope; diagnostic arthroscopy and management of intra-articular lesions; ATFL remnant assessment; anchor type, size, and number; fibular anchor insertion technique and intended target (e.g., relation to the FOT); suture passage and tensioning; IER (Gould) augmentation; CFL management.]]')
h('Postoperative Rehabilitation', 2)
para('[[To be provided: immobilization type and duration, weightbearing progression, range-of-motion and strengthening program, timing of return to sports.]]')

h('3D CT Measurement of Anchor Position', 2)
para('Postoperative CT was obtained at [[time point]] using [[scanner]], and 3D reconstruction was performed with [[software]]. Following the method of Lee et al <<lee2022>>, the reconstructed image was rotated to a true lateral view of the distal fibula, and the FAT, FOT, and anchor center were identified (Figure 2). The distance from the anchor center to the FOT was divided by the distance between the FAT and FOT and expressed as a percentage. On the basis of cadaveric studies of the fibular ATFL footprint <<burks1994,matsui2017,kakegawa2019,nakasa2021>>, Lee et al <<lee2022>> defined the lower quarter between the FOT and FAT as the anatomic position; accordingly, ankles were classified into the anatomic (<25%), subanatomic (25% to <50%), and nonanatomic (≥50%) groups. Measurements were performed by [[number, profession, blinding]], and the reliability was [[ICC / kappa]].')

h('Clinical Evaluation', 2)
para('Pain was evaluated using a 10-point VAS. Function was evaluated using the KP score <<karlsson1991>> and the AOFAS ankle-hindfoot score <<kitaoka1994>>, and the activity level was evaluated using the Tegner activity score <<tegner1985>>. These 4 scores were assessed preoperatively and at the final follow-up. At the final follow-up, the FFI <<budiman1991>>, FAOS <<roos2001>>, FAAM activities of daily living (ADL) and sports subscales <<martin2005>>, and EQ-5D-5L <<herdman2011>> were additionally assessed. {{For the FAOS, the 5 subscales (pain, symptoms, ADL, sport and recreation, and quality of life) and their mean were reported. The EQ-5D-5L was reported as the sum of the 5 dimension levels (range, 5-25; lower is better).}} [[confirm definitions; Korean versions and their references]] Generalized joint laxity was assessed with the Beighton score <<beighton>>. The final evaluation was performed [[in person / by telephone]] at a mean of [[ ]] months (range, [[ ]]) after surgery.')

h('Radiologic Evaluation', 2)
para('Stress radiographs were obtained preoperatively and at the final follow-up using [[device, load]]. Varus talar tilt (degrees) was measured on the anteroposterior stress view, and anterior talar translation (mm) was measured on the lateral stress view [[confirm method]]. Stress radiographs of the contralateral ankle were obtained at the same time points, and the side-to-side difference was calculated as the value of the operated ankle minus that of the contralateral ankle. Measurements were performed by [[ ]].')

h('Statistical Analysis', 2)
para('Continuous variables are presented as mean ± standard deviation. Comparisons among the 3 groups were performed using the Kruskal-Wallis test for continuous variables and the chi-square test for categorical variables. When the Kruskal-Wallis test was significant, post hoc pairwise comparisons were performed using the Mann-Whitney U test with a Bonferroni-adjusted significance level of P < .017, as in the study by Lee et al <<lee2022>>. Preoperative and final follow-up values were compared using the Wilcoxon signed-rank test. To evaluate anchor position as a continuous variable, Spearman rank correlation coefficients between the anchor position ratio and the outcome variables were calculated. Statistical significance was set at P < .05. {{Analyses in this draft were performed with Python (SciPy 1.17).}} [[replace with the software actually used]] [[power analysis]]')

# ================================================================= RESULTS
h('Results')
h('Patient Characteristics', 2)
para(f'A total of {N} ankles in {NPAT} patients ({male_pat} men and {fem_pat} women) were included. The mean age at surgery was {ms(S.age,1)} years (range, {S.age.min():.0f}-{S.age.max():.0f} years), and {right} right and {left} left ankles were operated on. The anchor position ratio was {ms(S.ratio,1)}% overall, {ms(S[S.G=="A"].ratio,1)}% in the anatomic group (n = {ng["A"]}, {pct(ng["A"],N)}%), {ms(S[S.G=="S"].ratio,1)}% in the subanatomic group (n = {ng["S"]}, {pct(ng["S"],N)}%), and {ms(S[S.G=="N"].ratio,1)}% in the nonanatomic group (n = {ng["N"]}, {pct(ng["N"],N)}%). There were no significant differences among the 3 groups in age, sex, side, body mass index, Beighton score, or preoperative clinical and radiologic parameters (Table 1).')

h('Clinical Outcomes', 2)
para(f'In the overall cohort, the VAS decreased from {ms(S.pre_vas)} preoperatively to {ms(S.vas)} at the final follow-up, the KP score improved from {ms(S.pre_kp)} to {ms(S.kp)}, and the AOFAS score improved from {ms(S.pre_aofas)} to {ms(S.aofas)} (all P < .001). These improvements were significant within each group (Table 2). The Tegner activity score increased from {{{{{ms(S.teg_pre)} to {ms(S.teg_post)} (P = {fp(wil(S,"teg_pre","teg_post"))})}}}}; the activity level improved in {teg_up} ankles, was maintained in {teg_same}, and decreased in {teg_down}, so that it was maintained or improved in {{{{{teg_keep} ankles ({pct(teg_keep,N)}%)}}}}. The proportion of ankles with a maintained or improved Tegner score did not differ among the groups (P = {fp(chi(S.assign(k=S.d_teg>=0),"k"))}).')
para(f'The final VAS, KP, AOFAS, and Tegner scores and their pre- to postoperative changes did not differ significantly among the 3 groups (all P > .05; Table 2). Among the outcomes assessed only at the final follow-up, the FFI differed significantly among the groups (P = {fp(ffi_kw)}). In the post hoc analysis, the FFI was lower (better) in the subanatomic group than in the anatomic group ({ms(S[S.G=="S"].ffi)} vs {ms(S[S.G=="A"].ffi)}; P = {fp(ffi_ph["AS"])}), whereas neither group differed from the nonanatomic group (P = {fp(ffi_ph["AN"])} and P = {fp(ffi_ph["SN"])}, respectively). {{{{[[see query 5 — the abstract states this difference was not significant after correction]]}}}} The FAOS (total and all 5 subscales), FAAM ADL and sports subscales, and EQ-5D-5L did not differ among the groups (Table 3).')

h('Radiologic Outcomes', 2)
para(f'Final stress radiographs were available for {len(XR)} ankles (anatomic, {nx["A"]}; subanatomic, {nx["S"]}; nonanatomic, {nx["N"]}). In these ankles, the talar tilt decreased from {ms(XR.pre_tt)}° to {ms(XR.tt)}° and the anterior translation decreased from {ms(XR.pre_at)} mm to {ms(XR["at"])} mm (both P < .001). The side-to-side difference in talar tilt decreased from {ms(XRc.ssd_pre_tt)}° to {ms(XRc.ssd_tt)}° (P {fp(wil(XRc,"ssd_pre_tt","ssd_tt"))}), and that in anterior translation decreased from {ms(XRc.ssd_pre_at)} mm to {ms(XRc.ssd_at)} mm (P = {fp(wil(XRc,"ssd_pre_at","ssd_at"))}; n = {len(XRc)}). The final talar tilt, anterior translation, side-to-side differences, and pre- to postoperative changes did not differ among the 3 groups (all P > .05; Table 4). Within-group comparison showed significant decreases in talar tilt in all 3 groups, whereas the decrease in anterior translation was significant in the anatomic and subanatomic groups but not in the nonanatomic group ({ms(XR[XR.G=="N"].pre_at)} to {ms(XR[XR.G=="N"]["at"])} mm; P = {fp(wil(XR[XR.G=="N"],"pre_at","at"))}; n = {nx["N"]}).')

h('Anchor Position as a Continuous Variable', 2)
r_sig = [(lab, *rho(S if v not in ('d_tt', 'ssd_tt', 'd_at', 'ssd_at', 'tt', 'at') else XR, v)) for lab, v in
         [('AOFAS change', 'd_aofas'), ('talar tilt change', 'd_tt'), ('final side-to-side difference in talar tilt', 'ssd_tt')]]
para(f'The anchor position ratio was not significantly correlated with the final VAS, KP, AOFAS, Tegner, FFI, FAOS, FAAM, or EQ-5D-5L scores (Table 5). Weak correlations were observed between the ratio and the change in AOFAS score (ρ = {r_sig[0][1]:.2f}, P = {fp(r_sig[0][2])}), the change in talar tilt (ρ = {r_sig[1][1]:.2f}, P = {fp(r_sig[1][2])}), and the final side-to-side difference in talar tilt (ρ = {r_sig[2][1]:.2f}, P = {fp(r_sig[2][2])}); these did not indicate a consistent disadvantage of higher anchor positions {{{{[[interpretation to confirm; these were not corrected for multiple testing]]}}}}.')

# ================================================================= DISCUSSION
h('Discussion')
para(f'The principal finding of this study was that clinical and radiologic outcomes after the arthroscopic modified Broström procedure did not differ significantly according to the fibular anchor position measured on 3D CT. Pain, KP, and AOFAS scores improved significantly in all 3 groups, the magnitude of improvement was similar among the groups, and the activity level was maintained or improved in {pct(teg_keep,N)}% of the ankles. Stress radiographic stability at the final follow-up and its change from the preoperative values were also comparable among the groups. Therefore, our hypothesis that outcomes would differ according to anchor position was not supported.')
para(f'These results differ from those of Lee et al <<lee2022>>, who reported inferior FAOS, Karlsson scores, and posturographic fall risk in the nonanatomic group after all-inside arthroscopic ATFL repair. The distribution of anchor positions was similar between the 2 studies (anatomic, subanatomic, and nonanatomic: {pct(ng["A"],N)}%, {pct(ng["S"],N)}%, and {pct(ng["N"],N)}% in our study vs 30.0%, 52.5%, and 17.5% in their study), and the same 3D CT classification was used, so differences in grouping are unlikely to explain the discrepancy. Several other factors may be considered. First, the surgical procedures differed. Lee et al repaired the ATFL with a single knotless anchor without additional augmentation <<lee2022,leeyang2021>>, whereas our patients underwent the arthroscopic modified Broström procedure [[with IER augmentation — confirm]]. {{{{It is possible that additional augmentation contributes to stability and function regardless of the exact anchor position; however, the present study was not designed to test this, and Lee et al <<lee2021>> found no additional benefit of IER augmentation after all-inside arthroscopic ATFL repair.}}}} Second, the postoperative scores in the nonanatomic group of Lee et al were considerably lower than those in our nonanatomic group (e.g., Karlsson score 45.4 ± 27.4 vs KP score {ms(S[S.G=="N"].kp,1)}), and their nonanatomic group did not show significant improvement in most subjective scores <<lee2022>>. In contrast, the scores of all 3 groups in our study were high at the final follow-up, which may reflect differences in patient populations, techniques, or rehabilitation, and may also indicate a ceiling effect that limits the ability to detect small between-group differences. Third, [[follow-up duration, surgeon experience/learning curve, patient activity level — to discuss after data confirmed]].')
para('Biomechanically, nonanatomic ATFL repair with a proximal fibular tunnel has been shown to alter ankle kinematics and laxity in cadavers <<shoji2019>>. In the present study, however, final talar tilt, anterior translation, and side-to-side differences did not differ among the groups, which is consistent with the stress radiographic findings of Lee et al, who also found no between-group differences in anterior translation or talar tilt <<lee2022>>. Although the reduction in anterior translation did not reach significance within the nonanatomic group, this group had the smallest number of ankles with final stress radiographs, and the between-group comparison of changes was not significant; this finding should therefore be interpreted with caution. Static stress radiographs may not capture rotational laxity, such as the internal rotation laxity described in the cadaveric study <<shoji2019>>, and dynamic or rotational evaluation would be needed to clarify whether anchor position affects these parameters in vivo.')
para(f'The FFI was the only outcome that differed among the groups. However, the difference was observed between the anatomic and subanatomic groups, with the subanatomic group showing the better score, and neither group differed from the nonanatomic group. In addition, the FFI was not correlated with the anchor position ratio when analyzed as a continuous variable (ρ = {ffi_rho[0]:.2f}, P = {fp(ffi_rho[1])}), and the mean values in all groups were low (range, {S.groupby("G").ffi.mean().min():.1f}-{S.groupby("G").ffi.mean().max():.1f}). Because a higher anchor position did not lead to worse FFI, this finding does not suggest a detrimental effect of nonanatomic anchor placement and may be attributable to multiple comparisons across many outcome measures. {{{{[[please confirm this interpretation; see query 5]]}}}}')
para(f'Identification of the fibular ATFL footprint during arthroscopy can be difficult because the footprint may not be fully visible from standard portals <<teramoto2018>>. In our series, {pct(ng["A"],N)}% of anchors were placed in the anatomic zone and {pct(ng["N"],N)}% in the nonanatomic zone, a distribution similar to that reported by Lee et al <<lee2022>>. The FOT has been proposed as a reliable bony landmark for identifying the ligament footprints <<matsui2017,nakasa2021>>, and efforts to place the anchor at the anatomic footprint remain reasonable. Nevertheless, our results suggest that, at least within the range of anchor positions observed in this study, a modest deviation of the anchor from the anatomic footprint may not substantially compromise clinical or radiologic outcomes at a minimum of 2 years after the arthroscopic modified Broström procedure.')

h('Limitations', 2)
para(f'This study has several limitations. First, it was a retrospective study with a relatively small number of ankles in the nonanatomic group (n = {ng["N"]}); therefore, the possibility of a type II error cannot be excluded [[power analysis]]. Second, the FFI, FAOS, FAAM, and EQ-5D-5L were assessed only at the final follow-up, so pre- to postoperative changes in these measures could not be evaluated. Third, final stress radiographs were not available for {N - len(XR)} ankles, and the contralateral final stress radiographs were not available for {len(XR) - len(XRc)} additional ankle(s). Fourth, one patient underwent bilateral surgery and both ankles were analyzed as independent observations. Fifth, the 3D CT classification used in this study was adopted from a single previous study <<lee2022>>, and its cutoff values have not been validated against clinical outcomes. Sixth, [[single-surgeon series / learning curve; no MRI or dynamic evaluation of the repaired ligament; ceiling effect of the scores]].')

h('Conclusion')
para('No significant differences in clinical or radiologic outcomes were found according to the fibular anchor position after the arthroscopic modified Broström procedure at a minimum 2-year follow-up. Pain and functional scores improved significantly in all groups, and the activity level was maintained or improved in most patients.')

# ================================================================= TABLES
page_break()
h('Tables')
G3 = ['Anatomic (n = %d)' % ng['A'], 'Subanatomic (n = %d)' % ng['S'], 'Nonanatomic (n = %d)' % ng['N']]


def grp(df, v, dec=2):
    return [ms(df[df.G == g][v], dec) for g in GROUPS]


# Table 1
rows = []
rows.append(['Age, y', ms(S.age, 1)] + grp(S, 'age', 1) + [fp(kw(S, 'age'))])
rows.append(['Sex, male/female, n'] + [f'{(x.sex=="M").sum()}/{(x.sex=="F").sum()}' for x in [S] + [S[S.G == g] for g in GROUPS]] + [fp(chi(S, 'sex'))])
rows.append(['Side, right/left, n'] + [f'{(x.side==1).sum()}/{(x.side==2).sum()}' for x in [S] + [S[S.G == g] for g in GROUPS]] + [fp(chi(S, 'side'))])
rows.append(['BMI, kg/m²', ms(S.bmi, 1)] + grp(S, 'bmi', 1) + [fp(kw(S, 'bmi'))])
rows.append(['Beighton score', ms(S.beighton, 1)] + grp(S, 'beighton', 1) + [fp(kw(S, 'beighton'))])
rows.append(['Anchor position ratio, %', ms(S.ratio, 1)] + grp(S, 'ratio', 1) + ['—'])
for lab, v in [('Preop VAS', 'pre_vas'), ('Preop KP score', 'pre_kp'), ('Preop AOFAS score', 'pre_aofas'), ('Preop Tegner score', 'teg_pre'),
               ('Preop talar tilt, °', 'pre_tt'), ('Preop anterior translation, mm', 'pre_at'),
               ('Preop SSD talar tilt, °', 'ssd_pre_tt'), ('Preop SSD anterior translation, mm', 'ssd_pre_at')]:
    rows.append([lab, ms(S[v])] + grp(S, v) + [fp(kw(S, v))])
table(['Variable', f'Total (n = {N})'] + G3 + ['P'], rows, [4.2, 2.5, 2.5, 2.5, 2.5, 1.3],
      title='Table 1. Patient characteristics and preoperative parameters',
      note='Values are mean ± SD unless otherwise indicated. P values from the Kruskal-Wallis test (continuous) or chi-square test (categorical). One patient underwent bilateral surgery (81 ankles in 80 patients); sex is counted per ankle. AOFAS, American Orthopaedic Foot and Ankle Society; BMI, body mass index; KP, Karlsson-Peterson; SSD, side-to-side difference (operated − contralateral); VAS, visual analog scale. [[add symptom duration, concomitant lesions if available]]')

# Table 2
rows = []
for lab, a, b, dv in [('VAS', 'pre_vas', 'vas', 'd_vas'), ('KP score', 'pre_kp', 'kp', 'd_kp'), ('AOFAS score', 'pre_aofas', 'aofas', 'd_aofas'), ('Tegner score', 'teg_pre', 'teg_post', 'd_teg')]:
    rows.append([f'**{lab}**', '', '', '', '', ''])
    rows.append(['  Preoperative', ms(S[a])] + grp(S, a) + [fp(kw(S, a))])
    rows.append(['  Final follow-up', ms(S[b])] + grp(S, b) + [fp(kw(S, b))])
    rows.append(['  Change', ms(S[dv])] + grp(S, dv) + [fp(kw(S, dv))])
    rows.append(['  P (pre vs final)', fp(wil(S, a, b))] + [fp(wil(S[S.G == g], a, b)) for g in GROUPS] + [''])
tg = lambda x: f'{int((x.d_teg>0).sum())}/{int((x.d_teg==0).sum())}/{int((x.d_teg<0).sum())}'
rows.append(['Tegner improved/maintained/decreased, n', tg(S)] + [tg(S[S.G == g]) for g in GROUPS] + [fp(chi(S.assign(k=np.sign(S.d_teg)), 'k'))])
rows.append(['Tegner maintained or improved, n (%)', f'{teg_keep} ({pct(teg_keep,N)})'] + [f'{int((S[S.G==g].d_teg>=0).sum())} ({pct(int((S[S.G==g].d_teg>=0).sum()), ng[g])})' for g in GROUPS] + [fp(chi(S.assign(k=S.d_teg >= 0), 'k'))])
table(['Variable', f'Total (n = {N})'] + G3 + ['P'], rows, [4.2, 2.5, 2.5, 2.5, 2.5, 1.3],
      title='Table 2. Clinical outcomes evaluated preoperatively and at final follow-up',
      note='Values are mean ± SD unless otherwise indicated. Between-group P values from the Kruskal-Wallis test (chi-square test for Tegner categories); pre vs final P values from the Wilcoxon signed-rank test. Change = final − preoperative.')

# Table 3
rows = []
for lab, v in [('FFI', 'ffi'), ('FAOS, mean of 5 subscales', 'faos'), ('  Pain', 'faos_pain'), ('  Symptoms', 'faos_sym'), ('  ADL', 'faos_adl'),
               ('  Sport/Recreation', 'faos_sport'), ('  Quality of life', 'faos_qol'), ('FAAM ADL', 'faam_adl'), ('FAAM Sports', 'faam_sport'), ('EQ-5D-5L (sum of levels)', 'eq5d')]:
    rows.append([lab, ms(S[v])] + grp(S, v) + [fp(kw(S, v))])
table(['Variable', f'Total (n = {N})'] + G3 + ['P'], rows, [4.2, 2.5, 2.5, 2.5, 2.5, 1.3],
      title='Table 3. Patient-reported outcomes at final follow-up',
      note=f'Values are mean ± SD. P values from the Kruskal-Wallis test. Post hoc Mann-Whitney U test for FFI: anatomic vs subanatomic, P = {fp(ffi_ph["AS"])}; anatomic vs nonanatomic, P = {fp(ffi_ph["AN"])}; subanatomic vs nonanatomic, P = {fp(ffi_ph["SN"])} (significance level P < .017). FAOS subscales were missing for one ankle in the anatomic group (n = 28). Lower FFI and EQ-5D-5L sum scores indicate better status. ADL, activities of daily living; FAAM, Foot and Ankle Ability Measure; FAOS, Foot and Ankle Outcome Score; FFI, Foot Function Index. {{{{[[FFI scale (0-100?) and EQ-5D-5L scoring to confirm]]}}}}')

# Table 4
G3x = ['Anatomic (n = %d)' % nx['A'], 'Subanatomic (n = %d)' % nx['S'], 'Nonanatomic (n = %d)' % nx['N']]
rows = []
for lab, a, b, dv in [('Talar tilt, °', 'pre_tt', 'tt', 'd_tt'), ('Anterior translation, mm', 'pre_at', 'at', 'd_at')]:
    rows.append([f'**{lab}**', '', '', '', '', ''])
    rows.append(['  Preoperative', ms(XR[a])] + grp(XR, a) + [fp(kw(XR, a))])
    rows.append(['  Final follow-up', ms(XR[b])] + grp(XR, b) + [fp(kw(XR, b))])
    rows.append(['  Change', ms(XR[dv])] + grp(XR, dv) + [fp(kw(XR, dv))])
    rows.append(['  P (pre vs final)', fp(wil(XR, a, b))] + [fp(wil(XR[XR.G == g], a, b)) for g in GROUPS] + [''])
for lab, a, b in [('SSD talar tilt, °', 'ssd_pre_tt', 'ssd_tt'), ('SSD anterior translation, mm', 'ssd_pre_at', 'ssd_at')]:
    rows.append([f'**{lab}**', '', '', '', '', ''])
    rows.append(['  Preoperative', ms(XRc[a])] + grp(XRc, a) + [fp(kw(XRc, a))])
    rows.append(['  Final follow-up', ms(XRc[b])] + grp(XRc, b) + [fp(kw(XRc, b))])
    rows.append(['  P (pre vs final)', fp(wil(XRc, a, b))] + [fp(wil(XRc[XRc.G == g], a, b)) for g in GROUPS] + [''])
table(['Variable', f'Total (n = {len(XR)})'] + G3x + ['P'], rows, [4.2, 2.5, 2.5, 2.5, 2.5, 1.3],
      title='Table 4. Stress radiographic outcomes',
      note=f'Values are mean ± SD. Only ankles with final stress radiographs were included (n = {len(XR)}). Side-to-side difference (SSD) = operated − contralateral; SSD rows include only ankles with contralateral final radiographs (n = {len(XRc)}; anatomic {int((XRc.G=="A").sum())}, subanatomic {int((XRc.G=="S").sum())}, nonanatomic {int((XRc.G=="N").sum())}). Between-group P values from the Kruskal-Wallis test; pre vs final P values from the Wilcoxon signed-rank test.')

# Table 5
rows = []
for lab, v, df in [('Final VAS', 'vas', S), ('Final KP score', 'kp', S), ('Final AOFAS score', 'aofas', S), ('Final Tegner score', 'teg_post', S),
                   ('Change in VAS', 'd_vas', S), ('Change in KP score', 'd_kp', S), ('Change in AOFAS score', 'd_aofas', S), ('Change in Tegner score', 'd_teg', S),
                   ('FFI', 'ffi', S), ('FAOS (mean)', 'faos', S), ('FAAM ADL', 'faam_adl', S), ('FAAM Sports', 'faam_sport', S), ('EQ-5D-5L', 'eq5d', S),
                   ('Final talar tilt', 'tt', XR), ('Final anterior translation', 'at', XR), ('Change in talar tilt', 'd_tt', XR), ('Change in anterior translation', 'd_at', XR),
                   ('Final SSD talar tilt', 'ssd_tt', XRc), ('Final SSD anterior translation', 'ssd_at', XRc)]:
    r, p, n = rho(df, v)
    rows.append([lab, str(n), f'{r:.2f}', fp(p)])
table(['Outcome', 'n', 'Spearman ρ', 'P'], rows, [6.5, 1.5, 2.5, 2],
      title='Table 5. Correlation between the anchor position ratio (continuous) and outcomes',
      note='A higher ratio indicates a more proximal (nonanatomic) anchor position. P values are not adjusted for multiple comparisons.')

# ================================================================= FIGURES
h('Figure Legends')
para('**Figure 1.** Flow diagram of patient selection. [[to be drawn after exclusion numbers are provided]]')
para('**Figure 2.** Measurement of the fibular anchor position on 3D CT. (A) The FAT, FOT, and anchor center are identified. (B) On the true lateral view of the reconstructed distal fibula, the ratio of the distance between the anchor center and the FOT to the distance between the FAT and the FOT is calculated. Anatomic, <25%; subanatomic, 25% to <50%; nonanatomic, ≥50%. [[representative images needed]]')
para('**Figure 3.** [[optional: representative 3D CT images of the 3 groups, or scatter plot of anchor ratio vs final outcome]]')

# ================================================================= REFERENCES
page_break()
h('References')
for i, k in enumerate(ORDER, 1):
    ok, txt = REFS[k]
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.first_line_indent = Cm(-0.8)
    r = p.add_run(f'{i}. {txt}')
    r.font.size = Pt(10)
    if not ok:
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        r2 = p.add_run('  [not verified against original — please provide PDF]')
        r2.font.size = Pt(8)
        r2.italic = True
unused = [k for k in REFS if k not in ORDER]
assert not unused, unused

# ================================================================= KOREAN ABSTRACT (revised)
page_break()
h('Appendix. 국문 초록 (수정안)')
para('관절경적 변형 브로스트롬 술식에서 비골 봉합나사 위치에 따른 임상 및 방사선학적 결과', bold=True)
para('서론 관절경적 변형 브로스트롬 술식에서 전거비인대의 해부학적 부착부를 고려한 봉합나사 삽입이 중요하다고 알려져 있다. 본 연구에서는 3차원 컴퓨터단층촬영(CT)을 이용하여 비골 봉합나사의 위치를 해부학적, 준해부학적 및 비해부학적으로 분류하고, 이에 따른 임상적 및 방사선학적 결과를 비교하고자 하였다.')
para(f'방법 관절경적 변형 브로스트롬 술식을 시행하고 2년 이상 추시한 {NPAT}명 {N}예를 대상으로 하였다. 봉합나사 중심에서 비골 은폐 결절(fibular obscure tubercle)까지의 거리를 비골 전방 결절(fibular anterior tubercle)과 비골 은폐 결절 사이 거리로 나눈 비율에 따라 25% 미만 값을 해부학적군({ng["A"]}예), 25% 이상 50% 미만 값을 준해부학적군({ng["S"]}예), 50% 이상 값을 비해부학적군({ng["N"]}예)으로 분류하였다. 시각적 통증 척도(Visual Analogue Scale, VAS), Karlsson-Peterson(KP), AOFAS 점수 및 Tegner 활동점수는 수술 전과 최종 추시 시점을 비교하였고, FFI, FAOS, EQ-5D-5L 및 FAAM은 최종 추시 시점 값을 분석하였다. 방사선학적으로 거골경사각, 전방 전위 및 건측과의 차이를 평가하였다.')
para(f'결과 VAS는 {S.pre_vas.mean():.2f}±{S.pre_vas.std():.2f}점에서 {S.vas.mean():.2f}±{S.vas.std():.2f}점으로 감소하였고 (P<.001), KP점수는 {S.pre_kp.mean():.2f}±{S.pre_kp.std():.2f}점에서 {S.kp.mean():.2f}±{S.kp.std():.2f}점, AOFAS 점수는 {S.pre_aofas.mean():.2f}±{S.pre_aofas.std():.2f}점에서 {S.aofas.mean():.2f}±{S.aofas.std():.2f}점으로 향상되었다 (P<.001). {{{{Tegner 활동점수는 {S.teg_pre.mean():.2f}±{S.teg_pre.std():.2f}점에서 {S.teg_post.mean():.2f}±{S.teg_post.std():.2f}점으로 증가하였으며 (P={fp(wil(S,"teg_pre","teg_post"))}), {teg_keep}예 ({pct(teg_keep,N)}%)에서 유지 또는 향상되었다.}}}} VAS, KP, AOFAS 및 Tegner 활동점수의 수술 전후 변화량도 세 군 간 차이가 없었다 (P>.05). {{{{FFI는 준해부학적군에서 해부학적군보다 낮았으나 (P={fp(ffi_kw)}), 봉합나사 위치 비율과의 상관관계는 없었다 (ρ={ffi_rho[0]:.2f}, P={fp(ffi_rho[1])}).}}}} 최종 거골경사각, 전방 전위, 건측과의 차이 및 수술 전후 방사선학적 변화량도 세 군 간 차이가 없었다 (P>.05).')
para('결론 비골 봉합나사 위치에 따른 임상적 및 방사선학적 결과의 유의한 차이는 확인되지 않았다. 전체 환자군에서 통증과 기능점수가 유의하게 호전되었으며, 활동 수준은 대부분 유지되거나 향상되었다.')

for p in doc.paragraphs:
    for r in p.runs:
        if 'P <.001' in r.text:
            r.text = r.text.replace('P <.001', 'P < .001')
doc.save(OUT)
print('saved', OUT, 'refs', len(ORDER))
