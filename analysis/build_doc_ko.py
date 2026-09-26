import re
import sys
from docx import Document
from docx.enum.text import WD_COLOR_INDEX, WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Pt, Cm
from docx.oxml.ns import qn
from stats_core import *

OUT = sys.argv[1]

# ---------------------------------------------------------------- references
# verified=True  -> citation copied from the reference list of Lee et al. (AJSM 2022) PDF provided by the author
# verified=False -> added from general knowledge; original must be checked by the author (highlighted in the list)
REFS = {
    'lee2022': (True, 'Lee SH, Yang JH, Kim I. Nonanatomic all-inside arthroscopic anterior talofibular ligament repair with a high-position anchor versus anatomic repair: an analysis based on 3D CT. Am J Sports Med. 2022;50(8):2134-2144. doi:10.1177/03635465221097119'),
    'brostrom': (True, 'Broström L. Sprained ankles: V. Treatment and prognosis in recent ligament ruptures. Acta Chir Scand. 1966;132:537-550.'),
    'gould': (False, 'Gould N, Seligson D, Gassman J. Early and late repair of lateral ligament of the ankle. Foot Ankle. 1980;1(2):84-89.'),
    'vega2013': (True, 'Vega J, Golanó P, Pellegrino A, Rabat E, Peña F. All-inside arthroscopic lateral collateral ligament repair for ankle instability with a knotless suture anchor technique. Foot Ankle Int. 2013;34:1701-1709.'),
    'takao2016': (True, 'Takao M, Matsui K, Stone JW, et al. Arthroscopic anterior talofibular ligament repair for lateral instability of the ankle. Knee Surg Sports Traumatol Arthrosc. 2016;24:1003-1006.'),
    'matsui2016': (True, 'Matsui K, Takao M, Miyamoto W, Matsushita T. Early recovery after arthroscopic repair compared to open repair of the anterior talofibular ligament for lateral instability of the ankle. Arch Orthop Trauma Surg. 2016;136:93-100.'),
    'matsui2014': (True, 'Matsui K, Takao M, Miyamoto W, Innami K, Matsushita T. Arthroscopic Broström repair with Gould augmentation via an accessory anterolateral port for lateral instability of the ankle. Arch Orthop Trauma Surg. 2014;134:1461-1467.'),
    'vega2020': (True, 'Vega J, Malagelada F, Dalmau-Pastor M. Arthroscopic all-inside ATFL and CFL repair is feasible and provides excellent results in patients with chronic ankle instability. Knee Surg Sports Traumatol Arthrosc. 2020;28:116-123.'),
    'drakos2014': (True, 'Drakos MC, Behrens SB, Paller D, Murphy C, DiGiovanni CW. Biomechanical comparison of an open vs arthroscopic approach for lateral ankle instability. Foot Ankle Int. 2014;35:809-815.'),
    'krips2000': (True, 'Krips R, van Dijk CN, Halasi T, et al. Anatomical reconstruction versus tenodesis for the treatment of chronic anterolateral instability of the ankle joint: a 2- to 10-year follow-up, multicenter study. Knee Surg Sports Traumatol Arthrosc. 2000;8:173-179.'),
    'shoji2019': (True, 'Shoji H, Teramoto A, Sakakibara Y, et al. Kinematics and laxity of the ankle joint in anatomic and nonanatomic anterior talofibular ligament repair: a biomechanical cadaveric study. Am J Sports Med. 2019;47:667-673.'),
    'caputo2009': (True, 'Caputo AM, Lee JY, Spritzer CE, et al. In vivo kinematics of the tibiotalar joint after lateral ankle instability. Am J Sports Med. 2009;37:2241-2248.'),
    'burks1994': (True, 'Burks RT, Morgan J. Anatomy of the lateral ankle ligaments. Am J Sports Med. 1994;22:72-77.'),
    'kakegawa2019': (True, 'Kakegawa A, Mori Y, Tsuchiya A, Sumitomo N, Fukushima N, Moriizumi T. Independent attachment of lateral ankle ligaments: anterior talofibular and calcaneofibular ligaments—a cadaveric study. J Foot Ankle Surg. 2019;58:717-722.'),
    'matsui2017': (True, 'Matsui K, Oliva XM, Takao M, et al. Bony landmarks available for minimally invasive lateral ankle stabilization surgery: a cadaveric anatomical study. Knee Surg Sports Traumatol Arthrosc. 2017;25:1916-1924.'),
    'nakasa2021': (True, 'Nakasa T, Ikuta Y, Ota Y, et al. Safe angles of ATFL and CFL anchor insertion into anatomical attachment of fibula in a lateral ankle ligament repair. J Orthop Sci. 2021;26:156-161. doi:10.1016/j.jos.2020.02.011'),
    'teramoto2018': (True, 'Teramoto A, Shoji H, Sakakibara Y, Suzuki T, Watanabe K, Yamashita T. The distal margin of the lateral malleolus visible under ankle arthroscopy (articular tip) from the anteromedial portal, is separate from the ATFL attachment site of the fibula: a cadaver study. J Orthop Sci. 2018;23:565-569.'),
    'thes2016': (True, 'Thès A, Klouche S, Ferrand M, Hardy P, Bauer T. Assessment of the feasibility of arthroscopic visualization of the lateral ligament of the ankle: a cadaveric study. Knee Surg Sports Traumatol Arthrosc. 2016;24:985-990.'),
    'lee2021': (True, 'Lee SH, Cho HG, Yang JH. Additional inferior extensor retinaculum augmentation after all-inside arthroscopic anterior talofibular ligament repair for chronic ankle instability is not necessary. Am J Sports Med. 2021;49:1721-1731.'),
    'leeyang2021': (True, 'Lee SH, Yang JH. All-inside arthroscopic anatomic anterior talofibular ligament repair for anterolateral ankle instability using a knotless suture anchor, allowing for tension adjustment. Arthrosc Tech. 2021;10:e925-e929. doi:10.1016/j.eats.2020.11.013'),
    'beighton': (True, 'Beighton P, Solomon L, Soskolne CL. Articular mobility in an African population. Ann Rheum Dis. 1973;32:413-418.'),
    'karlsson1991': (False, 'Karlsson J, Peterson L. Evaluation of ankle joint function: the use of a scoring scale. Foot. 1991;1(1):15-19.'),
    'kitaoka1994': (False, 'Kitaoka HB, Alexander IJ, Adelaar RS, Nunley JA, Myerson MS, Sanders M. Clinical rating systems for the ankle-hindfoot, midfoot, hallux, and lesser toes. Foot Ankle Int. 1994;15(7):349-353.'),
    'tegner1985': (False, 'Tegner Y, Lysholm J. Rating systems in the evaluation of knee ligament injuries. Clin Orthop Relat Res. 1985;(198):43-49.'),
    'budiman1991': (False, 'Budiman-Mak E, Conrad KJ, Roach KE. The Foot Function Index: a measure of foot pain and disability. J Clin Epidemiol. 1991;44(6):561-570.'),
    'roos2001': (True, 'Roos EM, Brandsson S, Karlsson J. Validation of the Foot and Ankle Outcome Score for ankle ligament reconstruction. Foot Ankle Int. 2001;22:788-794.'),
    'martin2005': (False, 'Martin RL, Irrgang JJ, Burdett RG, Conti SF, Van Swearingen JM. Evidence of validity for the Foot and Ankle Ability Measure (FAAM). Foot Ankle Int. 2005;26(11):968-983.'),
    'herdman2011': (False, 'Herdman M, Gudex C, Lloyd A, et al. Development and preliminary testing of the new five-level version of EQ-5D (EQ-5D-5L). Qual Life Res. 2011;20(10):1727-1736.'),
}
ORDER = []


def cite(keys):
    nums = []
    for k in keys.split(','):
        k = k.strip()
        if k not in REFS:
            raise KeyError(k)
        if k not in ORDER:
            ORDER.append(k)
        nums.append(ORDER.index(k) + 1)
    nums = sorted(set(nums))
    # compress ranges
    parts, i = [], 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        parts.append(str(nums[i]) if j == i else (f'{nums[i]},{nums[j]}' if j == i + 1 else f'{nums[i]}-{nums[j]}'))
        i = j + 1
    return '[' + ','.join(parts) + ']'


# ---------------------------------------------------------------- docx helpers
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
for side in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
    setattr(sec, side, Cm(2.5))
st = doc.styles['Normal']
st.font.name = 'Times New Roman'
st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn('w:eastAsia'), '바탕')
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.5
for lvl in (1, 2, 3):
    hs = doc.styles[f'Heading {lvl}']
    hs.font.name = 'Times New Roman'
    hs.element.rPr.rFonts.set(qn('w:eastAsia'), '바탕')
    hs.font.color.rgb = None
    hs.font.size = Pt({1: 14, 2: 12, 3: 11}[lvl])

TOKEN = re.compile(r'(\[\[.*?\]\]|\{\{.*?\}\}|<<[^>]+>>|\*\*.*?\*\*)')


def add_runs(p, text, size=None, bold=False, italic=False):
    """[[x]] = yellow (author input needed); {{x}} = turquoise (changed vs. abstract / check);
    <<key,key>> = citation; **x** = bold."""
    for part in TOKEN.split(text):
        if not part:
            continue
        hl = None
        b = bold
        if part.startswith('[[') and part.endswith(']]'):
            part, hl = '[' + part[2:-2] + ']', WD_COLOR_INDEX.YELLOW
        elif part.startswith('{{') and part.endswith('}}'):
            part, hl = part[2:-2].replace('[[', '[').replace(']]', ']'), WD_COLOR_INDEX.TURQUOISE
            part = re.sub(r'<<([^>]+)>>', lambda m: cite(m.group(1)), part)
        elif part.startswith('<<') and part.endswith('>>'):
            part = cite(part[2:-2])
        elif part.startswith('**') and part.endswith('**'):
            part, b = part[2:-2], True
        r = p.add_run(part)
        r.bold = b
        r.italic = italic
        if size:
            r.font.size = Pt(size)
        if hl:
            r.font.highlight_color = hl
    return p


def para(text, style=None, align=None, size=None, bold=False, italic=False, space_after=None, line=None):
    p = doc.add_paragraph(style=style)
    add_runs(p, text, size=size, bold=bold, italic=italic)
    if align:
        p.alignment = align
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    if line is not None:
        p.paragraph_format.line_spacing = line
    return p


def h(text, lvl=1):
    return doc.add_heading(text, level=lvl)


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def table(header, rows, widths, note=None, title=None):
    if title:
        para(title, bold=True, size=10, space_after=3, line=1.0)
    t = doc.add_table(rows=1, cols=len(header))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, txt in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ''
        add_runs(c.paragraphs[0], txt, size=8.5, bold=True)
    for row in rows:
        cells = t.add_row().cells
        for i, txt in enumerate(row):
            cells[i].text = ''
            add_runs(cells[i].paragraphs[0], str(txt), size=8.5, bold=(i == 0 and str(txt).isupper()))
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
            for p in row.cells[i].paragraphs:
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_after = Pt(0)
                if i > 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if note:
        para(note, size=8.5, line=1.0, space_after=12)
    return t


# ---------------------------------------------------------------- numbers
N = len(S)
NPAT = S.id.nunique()
ng = {g: int((S.G == g).sum()) for g in GROUPS}
nx = {g: int((XR.G == g).sum()) for g in GROUPS}
teg_up = int((S.d_teg > 0).sum())
teg_same = int((S.d_teg == 0).sum())
teg_down = int((S.d_teg < 0).sum())
teg_keep = teg_up + teg_same
pct = lambda a, b: f'{a / b * 100:.1f}'
ffi_kw = kw(S, 'ffi')
ffi_ph = mwu_pairs(S, 'ffi')
ffi_rho = rho(S, 'ffi')
fem_pat = int((S.drop_duplicates('id').sex == 'F').sum())
male_pat = NPAT - fem_pat
right = int((S.side == 1).sum())
left = int((S.side == 2).sum())
XRc = XR[XR.c_tt.notna()]
opmin, opmax = S.opdate.min(), S.opdate.max()

fmin = lambda v: f'{v:.2f}'
kp_n = ms(S[S.G == "N"].kp, 1)

# ================================================================= 저자 확인 요청 사항
para('관절경적 변형 브로스트롬 술식 – 비골 봉합나사 위치 연구: 원고 초안 v0.2 (국문)', bold=True, size=14, line=1.0)
para('2026-09-26 작성 · 공동저자 검토용 · 투고본 아님', italic=True, size=9, line=1.0)
para('하이라이트 표시 방법', bold=True, line=1.0)
para('[[노란색 = 저자만 알 수 있는 정보로, 임의로 채우지 않은 부분 (기입 요청)]]', size=10, line=1.0)
para('{{청록색 = 학회 초록과 달라진 값이거나, 초안 작성자의 해석으로 확인이 필요한 문장}}', size=10, line=1.0)
para('참고문헌 중 노란색 항목은 원문과 대조하지 않은 문헌입니다(원문 PDF 요청). 나머지는 Lee 등(AJSM 2022)의 참고문헌 목록에서 그대로 옮겼습니다. 표의 모든 수치는 “update ver2.5.xlsx”에서 스크립트로 직접 계산했으며, 손으로 옮겨 적은 값은 없습니다.', size=10, line=1.0)

para('저자 확인 요청 사항', bold=True, size=12, line=1.0)
queries = [
    ('대상 (81예)', '81예 = Sheet1에서 CT 봉합나사 비율이 있는 증례 중 2025년 이전 수술례(No.1–94 중 CT/설문 결측 13예 제외). 동일 등록번호 환자(No.72, No.80)가 양측 수술 → 80명 81예. 이 정의가 맞는지 확인 부탁드립니다. 전체 수술 환자 수와 제외 사유별 인원(흐름도용)도 필요합니다.'),
    ('최소 2년 추시', 'No.94(수술일 2024-12-31)는 2026-09-26 기준 약 21개월로 “최소 2년 추시” 조건을 만족하지 않습니다. 최종 설문 시점과 포함 여부를 알려주세요. No.93(2024-09-24)은 정확히 2년입니다.'),
    ('추시 기간', '엑셀의 “Last F/U”는 외래 방문일로 보여 최종 설문 시점으로 쓸 수 없습니다. 각 증례의 최종 설문 시행일(또는 평균/범위 추시 기간)과 설문 방법(외래/전화/우편)을 알려주세요.'),
    ('Tegner (변경됨)', f'현재 데이터: Tegner {ms(S.teg_pre)} → {ms(S.teg_post)}, Wilcoxon P = {fp(wil(S, "teg_pre", "teg_post"))}. 유지/향상 {teg_keep}예 ({pct(teg_keep, N)}%): 향상 {teg_up}, 유지 {teg_same}, 감소 {teg_down}. 초록의 “68예 (87.7%)”는 68/81 = 84.0%로 맞지 않아 71/81 = 87.7%가 맞는 것으로 보입니다.'),
    ('FFI 사후분석 (중요)', f'초록에는 “다중비교 보정 후 유의하지 않음”으로 되어 있으나, 현재 데이터로는 Kruskal–Wallis P = {fp(ffi_kw)}, 사후 Mann–Whitney 해부학적군 vs 준해부학적군 P = {fp(ffi_ph["AS"])}(Bonferroni 기준 P < .017에서도 유의), Dunn–Bonferroni 보정 P = .035로 여전히 유의합니다. 초록에서 사용한 보정 방법을 알려주시면 맞추겠습니다. 본문은 현재 계산값 그대로 기술했습니다.'),
    ('“3차원 CT 분석”의 의미', f'초록의 “3차원 CT 분석에서는 유의하지 않았다”를 봉합나사 위치 비율(연속변수)과 FFI의 Spearman 상관(ρ = {ffi_rho[0]:.2f}, P = {fp(ffi_rho[1])})으로 해석했습니다. 다른 분석이었다면 알려주세요.'),
    ('통계 방법/프로그램', 'Lee 등과 같은 비모수 방법(Kruskal–Wallis, Wilcoxon 부호순위, 사후 Mann–Whitney U [P < .017], 카이제곱 검정)과 Spearman 상관을 Python(SciPy 1.17)으로 계산했습니다. 초록 수치(VAS/KP/AOFAS)는 정확히 재현되었습니다. 실제 사용 프로그램(SPSS 등)과 검정법, 사전 검정력 분석 여부를 알려주세요.'),
    ('수술 방법', '마취/체위, 삽입구, 관절경, 봉합나사 종류·크기·개수(엑셀 메모에 “Arthrex anchor”, “IER reinforce” 언급), 전거비인대 봉합 방법, 하신전지대 보강 여부(전례 시행?), 종비인대 처치, 동반 술식, 집도의 수(단일 술자?).'),
    ('재활', '술후 고정 기간, 체중부하 시작 시점, 운동 복귀 시점 등.'),
    ('CT 측정', 'CT 촬영 시점(술후 며칠), 장비, 3차원 재구성 프로그램, 측정자(인원, 맹검 여부), 측정 신뢰도(ICC/kappa).'),
    ('스트레스 방사선', '장비(Telos 등)·하중(N)·측정 방법, 촬영 시점. 참고: 술전 전방 전위가 건측과 거의 차이가 없습니다(평균 0.52 mm). 최종 스트레스 사진이 없는 9예(각 군 2/2/5)가 있습니다.'),
    ('점수 정의', 'FAOS 총점 = 5개 하위척도 평균으로 계산되어 있습니다(원래 FAOS는 총점을 정의하지 않음). No.51은 총점만 있고 하위척도가 없습니다. EQ-5D-5L은 5–11 값으로 효용지수가 아닌 5개 차원 수준의 합(5 = 최상)으로 보입니다 → 한국형 가치세트로 지수 변환할지 결정 필요. 한국어판 도구 사용 여부와 인용문헌.'),
    ('포함/제외 기준, IRB', '포함/제외 기준, IRB 승인번호, 동의서 면제 여부, 연구기관명(경희대학교병원/강동경희대학교병원 등), 저자 영문 이름.'),
    ('투고 학술지', '투고 예정 학술지(분량·참고문헌 양식)를 알려주시면 형식을 맞추겠습니다. 현재는 번호식(Vancouver) 참고문헌으로 작성했습니다.'),
    ('필요 문헌 원문', '노란색 참고문헌의 원문 PDF, 그리고 가능하면 (1) Lee 등(2022) 이후 봉합나사 위치/비해부학적 봉합 관련 임상 연구, (2) 관절경적 변형 브로스트롬(하신전지대 보강 포함) 장기 추시 연구, (3) 한국어판 FAOS/FAAM/FFI/EQ-5D 타당도 논문.'),
]
table(['#', '항목', '내용'], [(str(i + 1), a, b) for i, (a, b) in enumerate(queries)], [0.8, 3.2, 12])
page_break()

# ================================================================= 표지
para('원저', italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)
para('관절경적 변형 브로스트롬 술식에서 비골 봉합나사 위치에 따른 임상 및 방사선학적 결과: 최소 2년 추시, 3차원 컴퓨터단층촬영 분석',
     bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER)
para('Clinical and Radiologic Outcomes of Arthroscopic Modified Broström Procedure According to Fibular Anchor Position: A 3D CT–Based Analysis With a Minimum 2-Year Follow-up',
     italic=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
para('이정현, 서동욱, 경민규, 정비오*', align=WD_ALIGN_PARAGRAPH.CENTER)
para('경희대학교 의과대학 정형외과학교실 [[병원명/주소 확인]]', align=WD_ALIGN_PARAGRAPH.CENTER, size=10)
para('*교신저자: 정비오 [[주소, 전화]], 이메일: biojeong@gmail.com', align=WD_ALIGN_PARAGRAPH.CENTER, size=10)
para('이해 상충: [[ ]]  ·  연구비: [[ ]]  ·  IRB 승인: [[기관, 번호]]', align=WD_ALIGN_PARAGRAPH.CENTER, size=10)
page_break()

# ================================================================= 초록
h('초록')
para('**목적:** 관절경적 외측 발목 인대 봉합술에서 비골 봉합나사를 해부학적 위치에 삽입하는 것이 중요하다고 알려져 있으며, 이전 3차원 컴퓨터단층촬영(computed tomography, CT) 연구에서는 전내시경적 전거비인대(anterior talofibular ligament, ATFL) 봉합술 후 봉합나사가 비해부학적(높은) 위치에 삽입된 경우 결과가 불량하다고 보고되었다. 본 연구에서는 관절경적 변형 브로스트롬 술식에서 비골 봉합나사 위치에 따라 임상 및 방사선학적 결과에 차이가 있는지 알아보고자 하였다.')
para(f'**대상 및 방법:** 관절경적 변형 브로스트롬 술식을 시행하고 2년 이상 추시한{{{{ [[확인 필요: 요청 사항 2]]}}}} {NPAT}명 {N}예를 후향적으로 분석하였다. 술후 3차원 CT에서 봉합나사 중심으로부터 비골 은폐 결절(fibular obscure tubercle, FOT)까지의 거리를 비골 전방 결절(fibular anterior tubercle, FAT)과 FOT 사이 거리로 나눈 비율에 따라 해부학적군(25% 미만, {ng["A"]}예), 준해부학적군(25% 이상 50% 미만, {ng["S"]}예), 비해부학적군(50% 이상, {ng["N"]}예)으로 분류하였다. 시각적 통증 척도(visual analogue scale, VAS), Karlsson-Peterson(KP) 점수, 미국족부족관절학회(American Orthopaedic Foot and Ankle Society, AOFAS) 점수 및 Tegner 활동 점수는 수술 전과 최종 추시 시점을 비교하였고, Foot Function Index(FFI), Foot and Ankle Outcome Score(FAOS), EQ-5D-5L 및 Foot and Ankle Ability Measure(FAAM)는 최종 추시 시점 값을 분석하였다. 방사선학적으로 스트레스 방사선 사진에서 거골경사각, 전방 전위 및 건측과의 차이를 평가하였다.')
para(f'**결과:** VAS는 {ms(S.pre_vas)}점에서 {ms(S.vas)}점으로 감소하였고(P < .001), KP 점수는 {ms(S.pre_kp)}점에서 {ms(S.kp)}점, AOFAS 점수는 {ms(S.pre_aofas)}점에서 {ms(S.aofas)}점으로 향상되었다(모두 P < .001). {{{{Tegner 활동 점수는 {ms(S.teg_pre)}점에서 {ms(S.teg_post)}점으로 증가하였으며(P = {fp(wil(S,"teg_pre","teg_post"))}), {teg_keep}예({pct(teg_keep,N)}%)에서 유지 또는 향상되었다.}}}} VAS, KP, AOFAS 및 Tegner 점수의 최종값과 수술 전후 변화량은 세 군 간 차이가 없었다(모두 P > .05). FFI는 군 간 차이를 보여(P = {fp(ffi_kw)}) 준해부학적군에서 해부학적군보다 낮았으나{{{{(사후분석 P = {fp(ffi_ph["AS"])}) [[요청 사항 5 참조]]}}}}, 봉합나사 위치 비율과의 상관관계는 없었다(ρ = {ffi_rho[0]:.2f}, P = {fp(ffi_rho[1])}). FAOS, FAAM 및 EQ-5D-5L은 군 간 차이가 없었다. 최종 거골경사각, 전방 전위, 건측과의 차이 및 수술 전후 방사선학적 변화량도 세 군 간 차이가 없었다(모두 P > .05).')
para('**결론:** 관절경적 변형 브로스트롬 술식 후 비골 봉합나사 위치에 따른 임상적 및 방사선학적 결과의 유의한 차이는 확인되지 않았다. 전체 환자군에서 통증과 기능 점수가 유의하게 호전되었으며, 활동 수준은 대부분 유지되거나 향상되었다.')
para('**색인단어:** 전거비인대, 만성 발목 불안정성, 관절경적 변형 브로스트롬 술식, 봉합나사 위치, 3차원 컴퓨터단층촬영')
page_break()

# ================================================================= 서론
h('서론')
para('Broström이 처음 기술한 외측 발목 인대의 직접 해부학적 봉합술은 만성 발목 불안정성(chronic ankle instability, CAI)의 수술적 치료로 가장 선호되는 방법이며<<brostrom,lee2022>>, 하신전지대(inferior extensor retinaculum, IER)를 이용한 보강이 이 술식의 변형으로 추가되었다<<gould,matsui2014>>. 관절경 술기의 발전과 함께 전내시경적 ATFL 봉합술과 관절경적 브로스트롬 술식이 개방적 술식의 효과적인 대안으로 자리잡았다<<vega2013,takao2016,matsui2016,vega2020>>. 생역학 연구에서 개방적 및 관절경적 브로스트롬 봉합술은 비슷한 강도를 보였으며<<drakos2014>>, 임상 연구에서도 관절경적 봉합술 후 빠른 회복과 양호한 결과가 보고되었다<<matsui2016,vega2020>>.')
para('정상 해부학의 복원은 외측 발목 인대 수술 후 좋은 예후를 위한 중요한 요소로 여겨진다<<krips2000,lee2022>>. Shoji 등<<shoji2019>>은 사체 연구에서 해부학적 부착부보다 근위부에 비골 골터널을 만들어 ATFL을 봉합한 경우(비해부학적 봉합) 정상 상태에 비해 내번 및 내회전 운동학과 내회전 이완이 증가하여, ATFL 결손 발목과 유사한 양상을 보인다고 하였다<<caputo2009,shoji2019>>. 그러나 관절경으로 비골의 ATFL 부착부를 확인하는 것은 항상 쉽지 않다. 한 사체 연구에서는 일반적인 관절경 삽입구로는 외과 첨부에서 근위부 7–10 mm 부위만 관찰되어, 의도치 않게 봉합나사가 근위부에 삽입될 가능성이 제기되었고<<teramoto2018>>, 다른 연구에서는 관절경으로 외측 인대와 부착부를 확인하는 것이 가능하다고 보고하였다<<thes2016>>. 최소 침습 수술 중 인대 부착부를 찾기 위한 골성 지표로 FOT와 FAT가 제시되었다<<matsui2017,nakasa2021>>.')
para('Lee 등<<lee2022>>은 술후 3차원 CT를 이용하여 전내시경적 ATFL 봉합술 후 비골 봉합나사의 위치를 FAT와 FOT 사이의 상대적 위치에 따라 분류하였고, 높은 위치에 봉합나사를 삽입한 비해부학적 봉합군이 해부학적 봉합군에 비해 FAOS, Karlsson 점수 및 자세 검사상 낙상 위험도가 불량하였으나 스트레스 방사선 결과에는 군 간 차이가 없었다고 보고하였다. 이 연구의 술식은 단일 무매듭 봉합나사를 이용한 ATFL 봉합술이었다<<lee2022,leeyang2021>>. 관절경적 변형 브로스트롬 술식[[하신전지대 보강 포함 여부 확인]]에서도, 그리고 더 다양한 환자 보고 결과 지표를 사용하였을 때에도 봉합나사 위치가 결과에 비슷한 영향을 미치는지는 아직 명확하지 않다.')
para('이에 본 연구에서는 Lee 등<<lee2022>>이 제시한 3차원 CT 분류를 이용하여 관절경적 변형 브로스트롬 술식 후 비골 봉합나사 위치에 따른 임상 및 방사선학적 결과를 비교하고자 하였다. 이전 보고와 같이 봉합나사 위치에 따라 결과에 차이가 있을 것이라는 가설을 세웠다.')

# ================================================================= 대상 및 방법
h('대상 및 방법')
h('연구 대상', 2)
para(f'본 후향적 연구는 [[기관]] 기관윤리심의위원회의 승인([[승인번호]])을 받았으며, 동의서는 [[면제 / 취득]]되었다. {opmin.year}년 {opmin.month}월부터 {opmax.year}년 {opmax.month}월까지 CAI로 관절경적 변형 브로스트롬 술식을 시행받은 환자[[단일 술자 여부 확인]]를 대상으로 하였다. 포함 기준은 (1) [[ ]]개월 이상의 보존적 치료에 실패한 CAI, (2) 관절경적 변형 브로스트롬 술식 시행, (3) 술후 3차원 CT 시행, (4) 최소 2년 이상 임상 추시가 가능한 경우였다. 제외 기준은 [[예: 동측 발목의 이전 수술, 관절염, 인대 재건술이나 절골술 등 동반 술식, 신경근육 질환 — 기입 요청]]이었다. 총 [[ ]]예 중 [[ ]]예가 제외되어([[사유]]), 최종 {NPAT}명 {N}예(1명은 양측 수술)를 분석하였다(그림 1).')
h('수술 방법', 2)
para('[[기입 요청: 마취와 체위; 삽입구와 관절경; 진단적 관절경과 관절 내 병변 처치; ATFL 잔여 조직 평가; 봉합나사 종류, 크기, 개수; 비골 봉합나사 삽입 방법과 목표 위치(예: FOT와의 관계); 봉합사 통과 및 긴장 조절; 하신전지대(Gould) 보강; 종비인대 처치]]')
h('술후 재활', 2)
para('[[기입 요청: 고정 방법과 기간, 체중부하 진행, 관절 운동 및 근력 강화 프로그램, 운동 복귀 시점]]')
h('3차원 CT를 이용한 봉합나사 위치 측정', 2)
para('술후 [[시점]]에 [[장비]]로 CT를 촬영하였고, [[프로그램]]으로 3차원 재구성을 시행하였다. Lee 등<<lee2022>>의 방법에 따라 재구성 영상을 원위 비골의 정측면상이 되도록 회전시킨 후 FAT, FOT 및 봉합나사 중심을 확인하였다(그림 2). 봉합나사 중심에서 FOT까지의 거리를 FAT와 FOT 사이 거리로 나누어 백분율로 표시하였다. 비골 ATFL 부착부에 대한 사체 연구<<burks1994,matsui2017,kakegawa2019,nakasa2021>>를 근거로 Lee 등<<lee2022>>은 FOT와 FAT 사이의 하방 1/4을 해부학적 위치로 정의하였으며, 이에 따라 해부학적군(25% 미만), 준해부학적군(25% 이상 50% 미만), 비해부학적군(50% 이상)으로 분류하였다. 측정은 [[측정자 수, 직종, 맹검 여부]]가 시행하였고 측정 신뢰도는 [[ICC / kappa]]였다.')
h('임상적 평가', 2)
para('통증은 10점 척도의 VAS로, 기능은 KP 점수<<karlsson1991>>와 AOFAS 발목-후족부 점수<<kitaoka1994>>로, 활동 수준은 Tegner 활동 점수<<tegner1985>>로 평가하였으며, 이 4가지 점수는 수술 전과 최종 추시 시점에 평가하였다. 최종 추시 시에는 FFI<<budiman1991>>, FAOS<<roos2001>>, FAAM 일상생활(activities of daily living, ADL) 및 스포츠 하위척도<<martin2005>>, EQ-5D-5L<<herdman2011>>을 추가로 평가하였다. {{FAOS는 5개 하위척도(통증, 증상, 일상생활, 스포츠/여가, 삶의 질)와 그 평균을 제시하였고, EQ-5D-5L은 5개 차원 수준의 합(범위 5–25, 낮을수록 양호)으로 제시하였다.}} [[정의 및 한국어판 사용 여부와 인용문헌 확인]] 전신 관절 이완성은 Beighton 점수<<beighton>>로 평가하였다. 최종 평가는 술후 평균 [[ ]]개월(범위, [[ ]])에 [[외래 / 전화]]로 시행하였다.')
h('방사선학적 평가', 2)
para('수술 전과 최종 추시 시 [[장비, 하중]]를 이용하여 스트레스 방사선 사진을 촬영하였다. 전후면 스트레스 사진에서 거골경사각(°)을, 측면 스트레스 사진에서 전방 전위(mm)를 측정하였다[[측정 방법 확인]]. 같은 시점에 건측 발목의 스트레스 사진도 촬영하였으며, 건측과의 차이는 환측 값에서 건측 값을 뺀 값으로 정의하였다. 측정은 [[ ]]가 시행하였다.')
h('통계 분석', 2)
para('연속변수는 평균 ± 표준편차로 제시하였다. 세 군 간 비교는 연속변수에 Kruskal-Wallis 검정을, 범주형 변수에 카이제곱 검정을 사용하였다. Kruskal-Wallis 검정이 유의한 경우 Lee 등<<lee2022>>과 같이 Mann-Whitney U 검정으로 사후 분석을 시행하고 Bonferroni 보정에 따라 유의수준을 P < .017로 설정하였다. 수술 전과 최종 추시 값은 Wilcoxon 부호순위 검정으로 비교하였다. 봉합나사 위치를 연속변수로 평가하기 위해 봉합나사 위치 비율과 결과 변수 간 Spearman 순위 상관계수를 구하였다. 통계적 유의수준은 P < .05로 하였다. {{본 초안의 분석은 Python(SciPy 1.17)으로 시행하였다.}} [[실제 사용한 프로그램으로 교체]] [[검정력 분석]]')

# ================================================================= 결과
h('결과')
h('환자 특성', 2)
para(f'총 {NPAT}명 {N}예(남자 {male_pat}명, 여자 {fem_pat}명)가 포함되었다. 수술 시 평균 나이는 {ms(S.age,1)}세(범위, {S.age.min():.0f}–{S.age.max():.0f}세)였고, 우측 {right}예, 좌측 {left}예였다. 봉합나사 위치 비율은 전체 {ms(S.ratio,1)}%, 해부학적군 {ms(S[S.G=="A"].ratio,1)}%({ng["A"]}예, {pct(ng["A"],N)}%), 준해부학적군 {ms(S[S.G=="S"].ratio,1)}%({ng["S"]}예, {pct(ng["S"],N)}%), 비해부학적군 {ms(S[S.G=="N"].ratio,1)}%({ng["N"]}예, {pct(ng["N"],N)}%)였다. 나이, 성별, 수술 부위, 체질량지수, Beighton 점수 및 술전 임상·방사선학적 지표는 세 군 간 유의한 차이가 없었다(표 1).')
h('임상적 결과', 2)
para(f'전체 환자군에서 VAS는 술전 {ms(S.pre_vas)}점에서 최종 추시 시 {ms(S.vas)}점으로 감소하였고, KP 점수는 {ms(S.pre_kp)}점에서 {ms(S.kp)}점으로, AOFAS 점수는 {ms(S.pre_aofas)}점에서 {ms(S.aofas)}점으로 향상되었다(모두 P < .001). 이러한 호전은 각 군 내에서도 모두 유의하였다(표 2). {{{{Tegner 활동 점수는 {ms(S.teg_pre)}점에서 {ms(S.teg_post)}점으로 증가하였으며(P = {fp(wil(S,"teg_pre","teg_post"))})}}}}, {teg_up}예에서 향상, {teg_same}예에서 유지, {teg_down}예에서 감소하여 {{{{{teg_keep}예({pct(teg_keep,N)}%)}}}}에서 유지 또는 향상되었다. Tegner 점수가 유지 또는 향상된 비율은 군 간 차이가 없었다(P = {fp(chi(S.assign(k=S.d_teg>=0),"k"))}).')
para(f'VAS, KP, AOFAS 및 Tegner 점수의 최종값과 수술 전후 변화량은 세 군 간 유의한 차이가 없었다(모두 P > .05, 표 2). 최종 추시 시에만 평가한 지표 중 FFI는 군 간 유의한 차이를 보였다(P = {fp(ffi_kw)}). 사후 분석에서 준해부학적군의 FFI가 해부학적군보다 낮았으나(양호; {ms(S[S.G=="S"].ffi)} vs {ms(S[S.G=="A"].ffi)}, P = {fp(ffi_ph["AS"])}), 두 군 모두 비해부학적군과는 차이가 없었다(각각 P = {fp(ffi_ph["AN"])}, P = {fp(ffi_ph["SN"])}). {{{{[[요청 사항 5 참조 — 초록에는 보정 후 유의하지 않다고 기술됨]]}}}} FAOS(총점 및 5개 하위척도), FAAM 일상생활 및 스포츠 하위척도, EQ-5D-5L은 군 간 차이가 없었다(표 3).')
h('방사선학적 결과', 2)
para(f'최종 스트레스 방사선 사진은 {len(XR)}예(해부학적군 {nx["A"]}예, 준해부학적군 {nx["S"]}예, 비해부학적군 {nx["N"]}예)에서 확보되었다. 거골경사각은 {ms(XR.pre_tt)}°에서 {ms(XR.tt)}°로, 전방 전위는 {ms(XR.pre_at)} mm에서 {ms(XR["at"])} mm로 감소하였다(모두 P < .001). 건측과의 거골경사각 차이는 {ms(XRc.ssd_pre_tt)}°에서 {ms(XRc.ssd_tt)}°로(P {fp(wil(XRc,"ssd_pre_tt","ssd_tt"))}), 전방 전위 차이는 {ms(XRc.ssd_pre_at)} mm에서 {ms(XRc.ssd_at)} mm로 감소하였다(P = {fp(wil(XRc,"ssd_pre_at","ssd_at"))}, {len(XRc)}예). 최종 거골경사각, 전방 전위, 건측과의 차이 및 수술 전후 변화량은 세 군 간 차이가 없었다(모두 P > .05, 표 4). 군 내 비교에서 거골경사각은 세 군 모두 유의하게 감소하였으나, 전방 전위는 해부학적군과 준해부학적군에서만 유의하게 감소하였고 비해부학적군에서는 유의하지 않았다({ms(XR[XR.G=="N"].pre_at)} mm → {ms(XR[XR.G=="N"]["at"])} mm, P = {fp(wil(XR[XR.G=="N"],"pre_at","at"))}, {nx["N"]}예).')
h('연속변수로서의 봉합나사 위치', 2)
r_sig = [(lab, *rho(df, v)) for lab, v, df in [('AOFAS 변화량', 'd_aofas', S), ('거골경사각 변화량', 'd_tt', XR), ('최종 건측과의 거골경사각 차이', 'ssd_tt', XRc)]]
para(f'봉합나사 위치 비율은 최종 VAS, KP, AOFAS, Tegner, FFI, FAOS, FAAM 및 EQ-5D-5L 점수와 유의한 상관관계가 없었다(표 5). 봉합나사 위치 비율과 AOFAS 변화량(ρ = {r_sig[0][1]:.2f}, P = {fp(r_sig[0][2])}), 거골경사각 변화량(ρ = {r_sig[1][1]:.2f}, P = {fp(r_sig[1][2])}), 최종 건측과의 거골경사각 차이(ρ = {r_sig[2][1]:.2f}, P = {fp(r_sig[2][2])}) 사이에 약한 상관관계가 관찰되었으나, 그 방향이 일정하지 않아 봉합나사가 높을수록 결과가 불량하다는 일관된 경향은 보이지 않았다{{{{[[해석 확인; 다중비교 보정 전 값]]}}}}.')

# ================================================================= 고찰
h('고찰')
para(f'본 연구의 주요 결과는 관절경적 변형 브로스트롬 술식 후 3차원 CT로 측정한 비골 봉합나사 위치에 따라 임상 및 방사선학적 결과에 유의한 차이가 없었다는 것이다. 세 군 모두에서 통증, KP 및 AOFAS 점수가 유의하게 호전되었고 호전 정도도 군 간 비슷하였으며, {pct(teg_keep,N)}%에서 활동 수준이 유지 또는 향상되었다. 최종 추시 시 스트레스 방사선상 안정성과 술전 대비 변화량도 군 간 비슷하였다. 따라서 봉합나사 위치에 따라 결과가 다를 것이라는 가설은 지지되지 않았다.')
para(f'이는 전내시경적 ATFL 봉합술 후 비해부학적군에서 FAOS, Karlsson 점수 및 자세 검사상 낙상 위험도가 불량하였다는 Lee 등<<lee2022>>의 결과와 다르다. 두 연구의 봉합나사 위치 분포는 비슷하였고(해부학적, 준해부학적, 비해부학적: 본 연구 {pct(ng["A"],N)}%, {pct(ng["S"],N)}%, {pct(ng["N"],N)}% vs Lee 등 30.0%, 52.5%, 17.5%), 같은 3차원 CT 분류를 사용하였으므로 군 분류의 차이로 이 차이를 설명하기는 어렵다. 몇 가지 요인을 고려할 수 있다. 첫째, 술식이 달랐다. Lee 등은 추가 보강 없이 단일 무매듭 봉합나사로 ATFL을 봉합하였으나<<lee2022,leeyang2021>>, 본 연구의 환자들은 관절경적 변형 브로스트롬 술식[[하신전지대 보강 여부 확인]]을 시행받았다. {{{{추가 보강이 봉합나사 위치와 관계없이 안정성과 기능에 기여하였을 가능성이 있으나, 본 연구는 이를 검증하도록 설계되지 않았으며, Lee 등<<lee2021>>은 전내시경적 ATFL 봉합술 후 하신전지대 보강의 추가 이득이 없다고 보고하였다.}}}} 둘째, Lee 등의 비해부학적군의 술후 점수는 본 연구의 비해부학적군보다 상당히 낮았으며(예: Karlsson 점수 45.4 ± 27.4 vs KP 점수 {kp_n}점), 해당 군에서는 대부분의 주관적 점수가 유의하게 호전되지 않았다<<lee2022>>. 반면 본 연구에서는 세 군 모두 최종 추시 시 점수가 높았는데, 이는 환자군, 술기 또는 재활의 차이를 반영할 수 있으며, 천장 효과로 인해 작은 군 간 차이를 검출하기 어려웠을 가능성도 있다. 셋째, [[추시 기간, 술자 경험/학습 곡선, 환자 활동 수준 — 자료 확인 후 기술]].')
para('생역학적으로 근위부 비골 골터널을 이용한 비해부학적 ATFL 봉합은 사체에서 발목 운동학과 이완성을 변화시키는 것으로 보고되었다<<shoji2019>>. 그러나 본 연구에서는 최종 거골경사각, 전방 전위 및 건측과의 차이가 군 간 차이가 없었으며, 이는 전방 전위와 거골경사각에 군 간 차이가 없었던 Lee 등의 스트레스 방사선 결과와도 일치한다<<lee2022>>. 비해부학적군에서 전방 전위의 감소가 군 내에서 유의하지 않았으나, 이 군은 최종 스트레스 사진이 있는 증례 수가 가장 적었고 변화량의 군 간 비교는 유의하지 않았으므로 해석에 주의가 필요하다. 정적 스트레스 방사선 사진은 사체 연구에서 기술된 내회전 이완성<<shoji2019>>과 같은 회전 불안정성을 반영하지 못할 수 있으므로, 봉합나사 위치가 생체 내에서 이러한 지표에 영향을 미치는지 밝히기 위해서는 동적 또는 회전 평가가 필요하다.')
para(f'FFI는 군 간 차이를 보인 유일한 지표였다. 그러나 차이는 해부학적군과 준해부학적군 사이에서 준해부학적군이 더 양호한 방향으로 나타났으며, 두 군 모두 비해부학적군과는 차이가 없었다. 또한 봉합나사 위치를 연속변수로 분석하였을 때 FFI와 상관관계가 없었고(ρ = {ffi_rho[0]:.2f}, P = {fp(ffi_rho[1])}), 모든 군에서 평균값이 낮았다(범위, {S.groupby("G").ffi.mean().min():.1f}–{S.groupby("G").ffi.mean().max():.1f}). 봉합나사가 높을수록 FFI가 나빠지지 않았으므로 이 결과가 비해부학적 봉합나사 위치의 불리한 영향을 시사한다고 보기는 어려우며, 여러 결과 지표에 대한 다중 비교의 영향일 수 있다. {{{{[[해석 확인 요청; 요청 사항 5 참조]]}}}}')
para(f'관절경 중 비골의 ATFL 부착부는 일반적인 삽입구로 완전히 관찰되지 않을 수 있어 확인이 어려울 수 있다<<teramoto2018>>. 본 연구에서 봉합나사의 {pct(ng["A"],N)}%가 해부학적 위치에, {pct(ng["N"],N)}%가 비해부학적 위치에 삽입되어 Lee 등<<lee2022>>과 비슷한 분포를 보였다. FOT는 인대 부착부를 확인하는 신뢰할 만한 골성 지표로 제시되었으며<<matsui2017,nakasa2021>>, 봉합나사를 해부학적 부착부에 삽입하려는 노력은 여전히 타당하다. 다만 본 연구 결과는 적어도 본 연구에서 관찰된 봉합나사 위치 범위 내에서는 해부학적 부착부로부터의 어느 정도의 편위가 관절경적 변형 브로스트롬 술식 후 최소 2년 시점의 임상 및 방사선학적 결과를 크게 저하시키지 않을 수 있음을 시사한다.')
h('연구의 제한점', 2)
para(f'본 연구는 몇 가지 제한점이 있다. 첫째, 후향적 연구이며 비해부학적군의 증례 수가 적어({ng["N"]}예) 제2종 오류의 가능성을 배제할 수 없다[[검정력 분석]]. 둘째, FFI, FAOS, FAAM 및 EQ-5D-5L은 최종 추시 시에만 평가하여 수술 전후 변화를 분석할 수 없었다. 셋째, {N - len(XR)}예에서 최종 스트레스 사진이 없었고, 추가로 {len(XR) - len(XRc)}예에서 건측 최종 스트레스 사진이 없었다. 넷째, 양측 수술을 받은 1명의 두 발목을 독립된 관측치로 분석하였다. 다섯째, 본 연구에서 사용한 3차원 CT 분류는 단일 선행 연구<<lee2022>>에서 도입된 것으로, 절단값이 임상 결과에 대해 검증되지 않았다. 여섯째, [[단일 술자 / 학습 곡선; 봉합 인대에 대한 자기공명영상이나 동적 평가 부재; 점수의 천장 효과]].')

h('결론')
para('관절경적 변형 브로스트롬 술식 후 최소 2년 추시에서 비골 봉합나사 위치에 따른 임상 및 방사선학적 결과의 유의한 차이는 확인되지 않았다. 모든 군에서 통증과 기능 점수가 유의하게 호전되었으며, 활동 수준은 대부분 유지되거나 향상되었다.')

# ================================================================= 표
page_break()
h('표')
G3 = [f'해부학적군 (n = {ng["A"]})', f'준해부학적군 (n = {ng["S"]})', f'비해부학적군 (n = {ng["N"]})']
W = [4.2, 2.5, 2.5, 2.5, 2.5, 1.3]


def grp(df, v, dec=2):
    return [ms(df[df.G == g][v], dec) for g in GROUPS]


rows = [['나이 (세)', ms(S.age, 1)] + grp(S, 'age', 1) + [fp(kw(S, 'age'))],
        ['성별, 남/여 (예)'] + [f'{(x.sex=="M").sum()}/{(x.sex=="F").sum()}' for x in [S] + [S[S.G == g] for g in GROUPS]] + [fp(chi(S, 'sex'))],
        ['수술 부위, 우/좌 (예)'] + [f'{(x.side==1).sum()}/{(x.side==2).sum()}' for x in [S] + [S[S.G == g] for g in GROUPS]] + [fp(chi(S, 'side'))],
        ['체질량지수 (kg/m²)', ms(S.bmi, 1)] + grp(S, 'bmi', 1) + [fp(kw(S, 'bmi'))],
        ['Beighton 점수', ms(S.beighton, 1)] + grp(S, 'beighton', 1) + [fp(kw(S, 'beighton'))],
        ['봉합나사 위치 비율 (%)', ms(S.ratio, 1)] + grp(S, 'ratio', 1) + ['—']]
for lab, v in [('술전 VAS', 'pre_vas'), ('술전 KP 점수', 'pre_kp'), ('술전 AOFAS 점수', 'pre_aofas'), ('술전 Tegner 점수', 'teg_pre'),
               ('술전 거골경사각 (°)', 'pre_tt'), ('술전 전방 전위 (mm)', 'pre_at'),
               ('술전 건측과의 거골경사각 차이 (°)', 'ssd_pre_tt'), ('술전 건측과의 전방 전위 차이 (mm)', 'ssd_pre_at')]:
    rows.append([lab, ms(S[v])] + grp(S, v) + [fp(kw(S, v))])
table(['변수', f'전체 (n = {N})'] + G3 + ['P'], rows, W, title='표 1. 환자 특성 및 술전 지표',
      note='값은 평균 ± 표준편차 또는 증례 수. P 값은 Kruskal-Wallis 검정(연속변수) 또는 카이제곱 검정(범주형 변수). 1명이 양측 수술을 받아(80명 81예) 성별은 발목 단위로 집계함. 건측과의 차이 = 환측 − 건측. AOFAS, American Orthopaedic Foot and Ankle Society; KP, Karlsson-Peterson; VAS, visual analogue scale. [[증상 기간, 동반 병변 자료가 있으면 추가]]')

rows = []
for lab, a, b, dv in [('VAS', 'pre_vas', 'vas', 'd_vas'), ('KP 점수', 'pre_kp', 'kp', 'd_kp'), ('AOFAS 점수', 'pre_aofas', 'aofas', 'd_aofas'), ('Tegner 점수', 'teg_pre', 'teg_post', 'd_teg')]:
    rows.append([f'**{lab}**', '', '', '', '', ''])
    rows.append(['  술전', ms(S[a])] + grp(S, a) + [fp(kw(S, a))])
    rows.append(['  최종 추시', ms(S[b])] + grp(S, b) + [fp(kw(S, b))])
    rows.append(['  변화량', ms(S[dv])] + grp(S, dv) + [fp(kw(S, dv))])
    rows.append(['  P (술전 vs 최종)', fp(wil(S, a, b))] + [fp(wil(S[S.G == g], a, b)) for g in GROUPS] + [''])
tg = lambda x: f'{int((x.d_teg>0).sum())}/{int((x.d_teg==0).sum())}/{int((x.d_teg<0).sum())}'
rows.append(['Tegner 향상/유지/감소 (예)', tg(S)] + [tg(S[S.G == g]) for g in GROUPS] + [fp(chi(S.assign(k=np.sign(S.d_teg)), 'k'))])
rows.append(['Tegner 유지 또는 향상, 예 (%)', f'{teg_keep} ({pct(teg_keep,N)})'] + [f'{int((S[S.G==g].d_teg>=0).sum())} ({pct(int((S[S.G==g].d_teg>=0).sum()), ng[g])})' for g in GROUPS] + [fp(chi(S.assign(k=S.d_teg >= 0), 'k'))])
table(['변수', f'전체 (n = {N})'] + G3 + ['P'], rows, W, title='표 2. 수술 전과 최종 추시 시 임상적 결과',
      note='값은 평균 ± 표준편차 또는 증례 수. 군 간 P 값은 Kruskal-Wallis 검정(Tegner 범주는 카이제곱 검정), 술전 vs 최종 P 값은 Wilcoxon 부호순위 검정. 변화량 = 최종 − 술전.')

rows = []
for lab, v in [('FFI', 'ffi'), ('FAOS (5개 하위척도 평균)', 'faos'), ('  통증', 'faos_pain'), ('  증상', 'faos_sym'), ('  일상생활', 'faos_adl'),
               ('  스포츠/여가', 'faos_sport'), ('  삶의 질', 'faos_qol'), ('FAAM 일상생활', 'faam_adl'), ('FAAM 스포츠', 'faam_sport'), ('EQ-5D-5L (수준 합)', 'eq5d')]:
    rows.append([lab, ms(S[v])] + grp(S, v) + [fp(kw(S, v))])
table(['변수', f'전체 (n = {N})'] + G3 + ['P'], rows, W, title='표 3. 최종 추시 시 환자 보고 결과',
      note=f'값은 평균 ± 표준편차. P 값은 Kruskal-Wallis 검정. FFI 사후 Mann-Whitney U 검정: 해부학적군 vs 준해부학적군 P = {fp(ffi_ph["AS"])}, 해부학적군 vs 비해부학적군 P = {fp(ffi_ph["AN"])}, 준해부학적군 vs 비해부학적군 P = {fp(ffi_ph["SN"])} (유의수준 P < .017). 해부학적군 1예에서 FAOS 하위척도 결측(n = 28). FFI와 EQ-5D-5L 수준 합은 낮을수록 양호. FAAM, Foot and Ankle Ability Measure; FAOS, Foot and Ankle Outcome Score; FFI, Foot Function Index. {{{{[[FFI 척도 범위(0–100?)와 EQ-5D-5L 산출 방식 확인]]}}}}')

G3x = [f'해부학적군 (n = {nx["A"]})', f'준해부학적군 (n = {nx["S"]})', f'비해부학적군 (n = {nx["N"]})']
rows = []
for lab, a, b, dv in [('거골경사각 (°)', 'pre_tt', 'tt', 'd_tt'), ('전방 전위 (mm)', 'pre_at', 'at', 'd_at')]:
    rows.append([f'**{lab}**', '', '', '', '', ''])
    rows.append(['  술전', ms(XR[a])] + grp(XR, a) + [fp(kw(XR, a))])
    rows.append(['  최종 추시', ms(XR[b])] + grp(XR, b) + [fp(kw(XR, b))])
    rows.append(['  변화량', ms(XR[dv])] + grp(XR, dv) + [fp(kw(XR, dv))])
    rows.append(['  P (술전 vs 최종)', fp(wil(XR, a, b))] + [fp(wil(XR[XR.G == g], a, b)) for g in GROUPS] + [''])
for lab, a, b in [('건측과의 거골경사각 차이 (°)', 'ssd_pre_tt', 'ssd_tt'), ('건측과의 전방 전위 차이 (mm)', 'ssd_pre_at', 'ssd_at')]:
    rows.append([f'**{lab}**', '', '', '', '', ''])
    rows.append(['  술전', ms(XRc[a])] + grp(XRc, a) + [fp(kw(XRc, a))])
    rows.append(['  최종 추시', ms(XRc[b])] + grp(XRc, b) + [fp(kw(XRc, b))])
    rows.append(['  P (술전 vs 최종)', fp(wil(XRc, a, b))] + [fp(wil(XRc[XRc.G == g], a, b)) for g in GROUPS] + [''])
table(['변수', f'전체 (n = {len(XR)})'] + G3x + ['P'], rows, W, title='표 4. 스트레스 방사선 결과',
      note=f'값은 평균 ± 표준편차. 최종 스트레스 사진이 있는 증례만 포함(n = {len(XR)}). 건측과의 차이 = 환측 − 건측; 해당 행은 건측 최종 사진이 있는 증례만 포함(n = {len(XRc)}; 해부학적 {int((XRc.G=="A").sum())}, 준해부학적 {int((XRc.G=="S").sum())}, 비해부학적 {int((XRc.G=="N").sum())}). 군 간 P 값은 Kruskal-Wallis 검정, 술전 vs 최종 P 값은 Wilcoxon 부호순위 검정.')

rows = []
for lab, v, df in [('최종 VAS', 'vas', S), ('최종 KP 점수', 'kp', S), ('최종 AOFAS 점수', 'aofas', S), ('최종 Tegner 점수', 'teg_post', S),
                   ('VAS 변화량', 'd_vas', S), ('KP 점수 변화량', 'd_kp', S), ('AOFAS 점수 변화량', 'd_aofas', S), ('Tegner 점수 변화량', 'd_teg', S),
                   ('FFI', 'ffi', S), ('FAOS (평균)', 'faos', S), ('FAAM 일상생활', 'faam_adl', S), ('FAAM 스포츠', 'faam_sport', S), ('EQ-5D-5L', 'eq5d', S),
                   ('최종 거골경사각', 'tt', XR), ('최종 전방 전위', 'at', XR), ('거골경사각 변화량', 'd_tt', XR), ('전방 전위 변화량', 'd_at', XR),
                   ('최종 건측과의 거골경사각 차이', 'ssd_tt', XRc), ('최종 건측과의 전방 전위 차이', 'ssd_at', XRc)]:
    r, p, n = rho(df, v)
    rows.append([lab, str(n), f'{r:.2f}', fp(p)])
table(['결과 변수', 'n', 'Spearman ρ', 'P'], rows, [6.5, 1.5, 2.5, 2],
      title='표 5. 봉합나사 위치 비율(연속변수)과 결과 변수 간 상관관계',
      note='비율이 높을수록 봉합나사가 근위부(비해부학적)에 위치함을 의미. P 값은 다중비교 보정 전 값.')

# ================================================================= 그림 설명
h('그림 설명')
para('**그림 1.** 환자 선정 흐름도. [[제외 인원 확인 후 작성]]')
para('**그림 2.** 3차원 CT에서 비골 봉합나사 위치 측정. (A) FAT, FOT 및 봉합나사 중심을 확인한다. (B) 재구성한 원위 비골의 정측면상에서 봉합나사 중심–FOT 거리를 FAT–FOT 거리로 나눈 비율을 구한다. 해부학적, 25% 미만; 준해부학적, 25% 이상 50% 미만; 비해부학적, 50% 이상. [[대표 영상 필요]]')
para('**그림 3.** [[선택: 세 군의 대표 3차원 CT 영상 또는 봉합나사 위치 비율과 최종 결과의 산점도]]')

# ================================================================= 참고문헌
page_break()
h('참고문헌')
for i, k in enumerate(ORDER, 1):
    ok, txt = REFS[k]
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.first_line_indent = Cm(-0.8)
    r = p.add_run(f'{i}. {txt}')
    r.font.size = Pt(10)
    if not ok:
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        r2 = p.add_run('  [원문 미확인 — PDF 요청]')
        r2.font.size = Pt(8)
        r2.italic = True
unused = [k for k in REFS if k not in ORDER]
assert not unused, unused

for p in doc.paragraphs:
    for r in p.runs:
        if 'P <.001' in r.text:
            r.text = r.text.replace('P <.001', 'P < .001')
doc.save(OUT)
print('saved', OUT, 'refs', len(ORDER))
